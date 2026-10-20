"""Trình bày phiên bản hiện tại, thay toàn bộ bảng điểm toàn trận của planner cũ."""

from copy import deepcopy
import json
from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]


def main():
    path=ROOT/'output/docx/bao_cao_hexudon.docx'
    doc=Document(path)
    data=json.loads((ROOT/'experiments/tour_api_evidence/summary.json').read_text(encoding='utf8'))
    peers=json.loads((ROOT/'experiments/tour_results/server_baseline.json').read_text(encoding='utf8'))
    pipeline=json.loads((ROOT/'experiments/tour_results/pipeline.json').read_text(encoding='utf8'))
    assert len(data)==8 and all(r['server_score'][:2]==r['ud_bound'] and not r['mismatch_count'] for r in data)
    frames=sum(r['checked_frames'] for r in data)
    values=sum(r['checked_values'] for r in data)
    fmt=lambda n:f'{n:,}'.replace(',','.')
    headings={p.text:p for p in doc.paragraphs if p.style.name.startswith('Heading')}
    def find(prefix):
        return next(p for p in doc.paragraphs if p.text.startswith(prefix) and not p.style.name.startswith('toc'))
    body_style=find('HEXUDON là bài toán').style
    heading_styles={n:next(p.style for p in doc.paragraphs if p.style.name==f'Heading {n}') for n in [1,2]}
    blocks=list(doc._element.body)
    def extract(start,end=None):
        first=blocks.index(find(start)._p)
        last=blocks.index(find(end)._p) if end else len(blocks)-1
        return [deepcopy(e) for e in blocks[first:last]]
    path_results=extract('6.1. So sánh năm','6.2.')
    references=extract('Tài liệu tham khảo','Phụ lục A.')
    appendix_a=extract('Phụ lục A.','Phụ lục B.')
    appendix_c=extract('Phụ lục C.')
    chapter4=find('4. Thuật toán')._p
    for e in blocks[blocks.index(chapter4):-1]:
        doc._element.body.remove(e)

    def font(p,size=12.5,bold=False):
        for run in p.runs:
            run.font.name='Times New Roman';run.font.size=Pt(size)
            run.font.color.rgb=RGBColor(0,0,0);run.font.bold=bold
    def p(text='',level=0,caption=False,command=False):
        result=doc.add_paragraph(text,style=heading_styles[level] if level else body_style)
        result.paragraph_format.line_spacing=1.3
        result.paragraph_format.space_after=Pt(6)
        result.paragraph_format.keep_together=True
        result.paragraph_format.keep_with_next=bool(level or caption)
        result.alignment=WD_ALIGN_PARAGRAPH.LEFT if level or caption or command else WD_ALIGN_PARAGRAPH.JUSTIFY
        font(result,11 if command else 12.5,bool(level or caption))
        return result
    def table(headers,rows,widths):
        total=(doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin)/360000
        widths=[w*total/sum(widths) for w in widths]
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
        for column,width in zip(t.columns,widths):column.width=Cm(width)
        for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
        for row in rows:
            for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width)
                for paragraph in cell.paragraphs:
                    paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.line_spacing=1.1
                    paragraph.paragraph_format.space_before=Pt(3)
                    paragraph.paragraph_format.space_after=Pt(3)
                    paragraph.paragraph_format.keep_with_next=i<len(t.rows)-1
                    font(paragraph,10.5,i==0)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEDED');cell._tc.get_or_add_tcPr().append(shade)
        return t
    def restore(elements):
        for element in elements:doc._element.body.insert(-1,element)
    def replace(prefix,text):
        paragraph=find(prefix);paragraph.text=text;font(paragraph)

    # Các chương đầu được giữ bố cục, chỉ sửa phạm vi và kết luận dẫn nhập.
    replace('Ở lớp tìm đường, nhóm đối chiếu',
        'Ở lớp tìm đường, nhóm đối chiếu DFS, BFS, Greedy Best-First, Dijkstra và A* trên trạng thái có tài nguyên. Phương pháp lập kế hoạch của nhóm biểu diễn tuyến cho cả đội, dùng Dijkstra tạo các chặng đi, chèn tham lam phân công điểm ghé và tìm kiếm cục bộ phá rồi dựng lại để cải thiện tuyến. Lịch hẹn tiếp tế và cách chuẩn bị vị trí cuối ngày giúp đội xe duy trì khả năng thu thập trong cả trận. Đây là một phương pháp kết hợp các thuật toán truyền thống, không sử dụng học máy trong luồng nộp chính.')
    replace('Ở đợt kiểm chứng mới, nhóm nộp',
        f'Nhóm đã nộp phương pháp này lên website BTC trên 8 bản đồ Practice khác nhau, dùng đội hình đã chọn bằng mô phỏng và ngân sách 5 giây lập kế hoạch mỗi ngày. Máy chủ nhận đủ 80 kế hoạch ngày; cả 8 bài đều đạt cận trên của số loại toàn trận U và tổng số loại theo ngày D. Nhóm đối chiếu {fmt(frames)} khung trạng thái cùng {fmt(values)} giá trị trong replay và không phát hiện sai khác ở các trường kiểm tra. Các kết quả về điểm toàn trận trong phần chính của báo cáo đều là điểm BTC của phiên bản hiện tại.')
    replace('Các cận trên được tính riêng',
        'Nhóm tính riêng cận trên cho từng thành phần điểm từ số loại udon, số ngày, lượng hàng và đội hình. Vì lời giải hợp lệ đã đạt cận U và D, hai thành phần này được xác định là tối ưu trên các bài đã nộp; số suất Q vẫn chưa có chứng minh tối ưu. So sánh với điểm máy chủ đã đọc của đội 57 cho thấy phương pháp đạt bộ điểm cao hơn trên cả 8 bài cùng input. Kết quả này chỉ áp dụng cho các lượt Practice được đối chiếu, không phải kết luận về mọi bản đồ hay mọi đội tham gia.')
    old=find('Nhóm thực hiện thêm 77 trận')
    old._p.getparent().remove(old._p)
    replace('Một tuyến ngắn chưa đủ',
        'Một tuyến ngắn chưa đủ để tạo kế hoạch thu thập tốt cho cả đội xe. Chương trình còn phải chọn điểm ghé, phối hợp các xe, duy trì nhiên liệu và chuẩn bị vị trí cho ngày sau. Nhóm chọn bài toán này để nghiên cứu cách kết hợp các thuật toán đồ thị, tham lam và tìm kiếm cục bộ thành một hệ thống giải quyết yêu cầu thực tế. Bộ mô phỏng và API BTC cung cấp cơ sở kiểm tra, còn các quyết định thuật toán là trọng tâm của báo cáo.')
    replace('Các câu hỏi chính là:',
        'Nhóm tập trung vào ba câu hỏi: các thuật toán tìm đường khác nhau như thế nào khi có chi phí và nhiên liệu; cách chèn điểm và cải thiện tuyến có thể phối hợp một đội xe ra sao; và kết quả do máy chủ BTC chấm đã đạt tới đâu so với các cận trên. Phần so sánh tìm đường được đánh giá trên truy vấn độc lập, còn điểm toàn trận thuộc phương pháp lập tuyến kết hợp. Báo cáo không dùng tính tối ưu của một đường đi để suy ra tính tối ưu của toàn bộ trận.')
    replace('Sản phẩm gồm năm thuật toán',
        'Sản phẩm gồm năm thuật toán tìm đường có xét tài nguyên, phương pháp lập tuyến cả đội bằng chèn tham lam và tìm kiếm cục bộ, bộ chọn vai trò xe, bộ tính cận, simulator và validator. Nhóm cung cấp thêm script lấy input và nộp từng ngày, công cụ đối chiếu replay, số liệu tổng hợp từ máy chủ và hướng dẫn tái hiện. Các mô-đun thử nghiệm khác trong mã nguồn không được mặc định là thành phần của phương pháp được chấm trong báo cáo.')
    for paragraph in doc.paragraphs:
        if 'Chúng em chọn A* làm bộ tìm đường chung' in paragraph.text:
            paragraph.text=re.sub(r'Chúng em chọn A\* làm bộ tìm đường chung.*',
                'A* được dùng trong phép so sánh tìm đường có tài nguyên. Trong phương pháp lập tuyến cả đội, chương trình dùng Dijkstra với nhiều trọng số để tái sử dụng cây đường đi; cả hai lớp đều không tự bảo đảm điểm toàn trận tối ưu.',paragraph.text)
            font(paragraph)
    # Bảng 1 giới thiệu đúng các lớp thuật toán của phiên bản được báo cáo.
    intro=doc.tables[0]
    intro.cell(2,1).text='Chèn tham lam và tìm kiếm cục bộ phá rồi dựng lại'
    intro.cell(2,2).text='Điểm BTC (U, D, Q), cận trên và tính hợp lệ'
    intro.cell(4,0).text='Kiểm tra giới hạn của lời giải'
    intro.cell(4,1).text='Cận nguồn cung và Dijkstra theo nhiên liệu'
    intro.cell(4,2).text='Đạt cận U, D; khoảng cách với cận Q'
    for row in intro.rows[1:]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:font(paragraph,10.5)

    p('4. Phương pháp lập tuyến cả đội có ràng buộc tài nguyên',1)
    p('Chương trình biểu diễn kế hoạch một ngày bằng các tuyến của xe tuần tra và các điểm hẹn của xe tiếp tế. Mỗi tuyến là một dãy điểm udon; đường giữa hai điểm được tính trên bản đồ và giao thông của ngày hiện tại. Bộ lập kế hoạch đánh giá tập tuyến của cả đội, thay vì tối ưu một xe rồi bỏ qua tác động tới các xe còn lại. Trong mã nguồn, phương pháp này được gọi là tour.')
    p('Bảng 7. Vai trò của các thuật toán trong phương pháp lập tuyến cả đội',caption=True)
    table(['Thành phần','Thuật toán hoặc cách xử lý','Quyết định được giải quyết'],[
        ['Tạo các chặng đi','Dijkstra theo ba trọng số step/nhiên liệu','Các phương án đường giữa các điểm'],
        ['Dựng tuyến','Chèn tham lam trên tập tuyến cả đội','Xe nào ghé điểm nào, chèn ở đâu'],
        ['Cải thiện tuyến','Tìm kiếm cục bộ phá rồi dựng lại','Đổi phân công và thứ tự ghé điểm'],
        ['Duy trì hoạt động','Chia nhiên liệu, lịch hẹn, ước lượng ngày sau','Tiếp tế và vị trí dừng cuối ngày'],
        ['Chọn vai trò','Mô phỏng các ứng viên bằng nhiều tiến trình','Xe nào tuần tra hoặc tiếp tế'],
        ['Kiểm tra trước nộp','Simulator thực thi đồng thời','Lệnh, step, nhiên liệu và kết quả dự kiến']],[3.5,6,6.5])
    p('4.1. Dijkstra tạo chặng đi theo step và nhiên liệu',2)
    p('Chương trình dùng Dijkstra để tạo cây đường ngắn nhất theo trọng số a×step + b×nhiên liệu. Ba cặp (a,b) là (1; 0,001), (1; 1) và (0,001; 1), lần lượt thiên về đi nhanh, cân bằng và tiết kiệm nhiên liệu. Chi phí được tính theo ô nguồn và trạng thái đường thật của ngày. Cây đường đi được tính khi cần rồi lưu lại, giúp nhiều phép chèn sử dụng chung kết quả thay vì tìm lại từ đầu.')
    p('Một trọng số chỉ chọn ra một phương án đại diện, nên ba kiểu đường không bao phủ mọi khả năng trao đổi giữa step và nhiên liệu. Vì vậy, chương trình vẫn chạy thử toàn tuyến để kiểm tra tài nguyên khi ghép các chặng. Các điểm udon nằm trên đường đi cũng được tính vào tập đã ghé. Cách dùng Dijkstra này khác với A* có nhiều nhãn tài nguyên ở chương 3: mục tiêu ở đây là tạo chặng dùng lại được trong nhiều lần đánh giá tuyến.')
    p('4.2. Chèn tham lam và phối hợp kho giữa các xe',2)
    p('Ban đầu, mỗi xe tuần tra có một tuyến rỗng. Chương trình thử chèn các điểm bán vào một số vị trí có triển vọng trên tuyến của từng xe, rồi chọn phép chèn tốt nhất của cả đội. Mức ưu tiên lần lượt là loại mới toàn trận, loại mới trong ngày, suất còn có thể lấy và mức tăng chi phí nhỏ hơn. Sau mỗi lượt nhận, tập điểm ghé và lượng hàng dự kiến được cập nhật chung để các xe không cùng nhắm những suất không còn khả dụng.')
    p('Giả sử tuyến đang có đoạn a → b và cần ghé thêm p. Phần chi phí tăng được ước lượng bởi c(a,p) + c(p,b) − c(a,b). Các vị trí được sàng lọc theo chi phí rồi kiểm tra lại tuyến với quỹ step, nhiên liệu, đường đi qua và lịch hẹn tiếp tế. Khi nhiều phép chèn có lợi ích ngang nhau, phép tăng chi phí ít hơn thường để lại nhiều tài nguyên cho các điểm sau. Do không xét mọi tổ hợp và mọi vị trí chèn, phương pháp không bảo đảm tối ưu toàn cục.')
    p('Nếu C(p) là số xe tuần tra có ghé điểm p trong ngày, số suất dự kiến tại đó là min(stocks(p), C(p)). Mỗi xe chỉ nhận tại một điểm một lần trong ngày. Chương trình dùng tập loại của các điểm đã ghé để tính độ phủ ngày và loại mới toàn trận. Cuối cùng, validator chạy các lệnh đồng thời để kiểm tra kết quả theo timeline của trò chơi; bước này giúp phát hiện những sai lệch mà phép tính xấp xỉ khi dựng tuyến chưa thể hiện hết.')
    p('4.3. Tìm kiếm cục bộ phá rồi dựng lại',2)
    p('Sau khi có lời giải khả thi, chương trình dùng thời gian còn lại để thử các phương án lân cận. Một lượt có thể bỏ vài điểm, bỏ một đoạn tuyến hoặc đổi kiểu đường của một xe. Chương trình dựng lại phần bị bỏ bằng chèn tham lam có nhiễu nhỏ ở chi phí, nhờ đó có cơ hội đổi thứ tự và phân công xe thay vì lặp đúng lời giải trước.')
    p('Với β = 0 như cấu hình được đánh giá, bộ giá trị của một ngày gồm loại mới toàn trận, số loại trong ngày, số suất, âm nhiên liệu đã dùng và âm step. Chương trình so từ trái sang phải và nhận phương án không kém lời giải hiện tại. Hai thành phần tài nguyên chỉ phá hòa sau các tiêu chí thu thập. Nhờ luôn giữ một lời giải đang có, tìm kiếm có thể kết thúc theo ngân sách mà vẫn trả ra kế hoạch; đây không phải bằng chứng đã tìm hết không gian lời giải.')
    p('4.4. Phân bổ nhiên liệu và lịch hẹn tiếp tế',2)
    p('Nếu đội không có xe tiếp tế, chương trình giới hạn nhiên liệu dùng hôm nay theo tỷ lệ quỹ step của ngày trên tổng quỹ step còn lại; ở ngày cuối, xe có thể dùng hết nhiên liệu còn lại. Đây là một quy tắc phân bổ nhằm tránh việc các xe hoạt động mạnh ở đầu trận rồi không còn khả năng di chuyển về sau, không phải nghiệm tối ưu của bài toán phân phối nhiên liệu.')
    p('Nếu có xe tiếp tế, chương trình thử một số điểm hẹn: vị trí đang đứng, điểm trung tâm, nơi gần nhóm tuần tra hoặc nơi thuận lợi cho xe sắp cạn nhiên liệu. Xe tiếp tế đi đến điểm được chọn rồi chờ. Khi cần, xe tuần tra rẽ vào điểm hẹn và ở chung đủ thời gian để được nạp. Các tuyến còn được kiểm tra khả năng quay về điểm hẹn theo nhiên liệu. Trong phiên bản được chấm, mỗi xe tiếp tế chỉ có một điểm hẹn trong một ngày.')
    p('4.5. Phủ loại và chuẩn bị vị trí cho ngày tiếp theo',2)
    p('D có ưu tiên cao hơn Q, nên mất một loại trong một ngày không thể được bù bằng nhiều suất hơn khi U bằng nhau. Nếu lời giải đang thiếu loại, hàm cover_types dựng lại nhiều khung tuyến chỉ nhằm phủ loại, rồi chèn thêm suất vào khung tốt nhất. Nhiễu nhỏ ở chi phí giúp các lần dựng khảo sát những cách phân công khác nhau, đặc biệt khi ngày ngắn khiến bài toán phủ loại chặt hơn.')
    p('Khi còn ngày sau, chương trình dùng vị trí và nhiên liệu cuối ngày để ước lượng số loại có thể thu trong ngày kế tiếp. Nếu ước lượng còn thiếu, hàm find_anchors tìm các ô neo thuận lợi rồi dựng lại tuyến hôm nay với yêu cầu dừng ở các ô đó. Phương án neo được xét theo loại mới toàn trận, loại hôm nay và độ phủ dự đoán ngày mai, nên có thể giảm Q hôm nay để bảo vệ D ngày sau. Giao thông hôm nay được dùng làm dự báo trong phép ước lượng; dữ liệu ngày mai thật vẫn phải được lấy lại từ API.')
    p('4.6. Chọn vai trò xe bằng mô phỏng',2)
    p('Vị trí xuất phát được giữ đúng thứ tự input; chương trình chỉ chọn vai trò tuần tra hoặc tiếp tế. Bộ chọn tạo các đội hình ứng viên gồm toàn tuần tra, từng xe tiếp tế và một số cấu hình nhiều xe tiếp tế. Các ứng viên được mô phỏng bằng nhiều tiến trình, sàng lọc rồi đánh giá lại tối đa bốn đội hình tốt khi còn thời gian. Khi có lượt hoàn tất, bộ chọn so điểm (U,D,Q); nếu không có lượt hoàn tất, cách chọn theo số ngày đã chạy và điểm hiện có chỉ là ước lượng.')
    p('Đội hình dùng trong các bảng kết quả đã được chọn offline với ngân sách 20 giây và được cố định trước khi nộp lên BTC. Cả tám đội hình có một xe tiếp tế, nhưng chỉ số xe làm tiếp tế khác nhau tùy bản đồ. Vì vậy, các điểm máy chủ xác nhận khả năng của phương pháp với những đội hình đã lưu; chúng chưa chứng minh bộ chọn luôn tìm ra đội hình tốt nhất hoặc luôn hoàn tất dưới đồng hồ chọn xe trong trận thi thật.')
    p('4.7. Ngân sách tính toán và kiểm tra kế hoạch',2)
    p('Mỗi ngày, runner nhận trạng thái thật rồi gọi planner với ngân sách 5 giây trong đợt đánh giá. Planner dành phần lớn thời gian cho dựng và sửa tuyến, sau đó simulator kiểm tra output. Nếu không còn step để tiếp tục, xe chờ hết ngày; nếu không đủ step cho một di chuyển thì không được phát lệnh đó. Khi tour không trả được kế hoạch hợp lệ, hệ thống có nhánh dự phòng Greedy; số ngày sử dụng dự phòng được báo riêng, không gán chúng thành thành tích của tour.')
    p('Giả sử cây đường được tạo từ R ô nguồn, đồ thị có V ô và E cạnh, phần Dijkstra cho ba kiểu trọng số có chi phí gần O(3R(V+E)log V) đối với các cây thực sự được tính. Phần chèn và sửa còn phụ thuộc số xe, điểm bán, vị trí chèn và vòng cải thiện. Giới hạn theo đồng hồ là giới hạn mềm vì một số bước dựng hoặc ước lượng không bị ngắt giữa chừng. Vì thế, thời gian đo thực tế và khả năng nộp trong hạn cần được báo riêng.')
    p('4.8. Cận trên để đánh giá mức đạt được',2)
    p('Gọi B là số loại brand khác nhau, H là số ngày, K(p) là kho mỗi ngày tại điểm p và A là số xe tuần tra. Cận theo nguồn cung là U ≤ B, D ≤ H×B và Q ≤ H×ΣK(p). Với đội hình cố định, Q còn bị chặn bởi H×Σmin(K(p),A), vì mỗi xe chỉ nhận ở một điểm tối đa một lần trong ngày. Các cận Q này bỏ qua lộ trình, nhiên liệu và quỹ step, nên không có bảo đảm rằng tồn tại phương án đạt được.')
    p('Mã nguồn còn có bộ cận tiếp cận bằng Dijkstra theo nhiên liệu. Nó xét vị trí đầu, nhiên liệu dùng chung qua các ngày và khả năng tiếp tế được nới lỏng để loại một số điểm chắc chắn không tới được. Cận đó là điều kiện cần, không phải một kế hoạch thực thi. Trong báo cáo, khi điểm hợp lệ của BTC đạt cận U và D, hai thành phần được xác định là tối ưu trên bài tương ứng; Q/cận sơ bộ chỉ là mức thu so với giới hạn nới lỏng, không phải tỷ lệ so với nghiệm tối ưu đã biết.')

    p('5. Thiết kế đánh giá và kiểm chứng',1)
    p('5.1. Dữ liệu và phạm vi từng phép đo',2)
    p('Báo cáo tách phép đo tìm đường khỏi phép đo điểm toàn trận. Nhóm giữ nguyên các truy vấn và số liệu so sánh năm thuật toán ở chương 3. Phần điểm trận đấu chỉ sử dụng các lượt nộp của phương pháp tour hiện tại đã được BTC chấm. Không lấy kết quả của phương pháp này để gán cho từng thuật toán tìm đường hoặc cho một chế độ planner khác.')
    p('Bảng 8. Dữ liệu và kết quả được sử dụng trong báo cáo',caption=True)
    table(['Nhóm đánh giá','Dữ liệu','Kết quả được đo'],[
        ['Tìm đường','30 bản đồ 8×8, 16×16, 32×32; 90 truy vấn; 5 thuật toán','450 lượt chạy; khả năng tìm đường, step, nhiên liệu, nhãn và thời gian'],
        ['Lập tuyến cả đội','8 bản đồ Practice khác nhau, mỗi bài 10 ngày','80 kế hoạch được BTC nhận; điểm cuối trận, độ phủ từng ngày'],
        ['Kiểm chứng simulator','Output đã nộp và replay của chính 8 trận','Đối chiếu vị trí, nhiên liệu, thu hàng và giao thông'],
        ['Mốc tham khảo bên ngoài','Điểm máy chủ đội 57 trên cùng 8 input','So bộ điểm; không biết thuật toán hoặc ngân sách của đội đó']],[3.4,6.4,6.2])
    p('Các bài Practice là Q01, Q03, Q04, Q06, Q07, Q08, Q09 và Q10. Q02 trùng cấu hình Q01 nên không được tính thêm như một bản đồ độc lập; Q05 chưa có trong danh sách ID đã khảo sát. Bộ Practice này có kích thước 30×30 hoặc 32×32. Phép đo tìm đường trên bản đồ tự sinh giúp khảo sát thêm cỡ nhỏ, nhưng không thay thế thử nghiệm toàn trận trên mọi kích thước của đề.')
    p('5.2. Quy tắc so sánh và thời gian',2)
    p('Trong phép đo tìm đường, các thuật toán nhận cùng input, trạng thái đường, nhiên liệu và quỹ step; ngân sách là 0,25 giây cho một truy vấn. Kết quả đọc đồng thời khả năng thành công, dừng vì hết giờ và chất lượng tuyến. Đối với chi phí đường, báo cáo dùng tập truy vấn mà cả năm thuật toán cùng tìm được đường, đồng thời vẫn giữ bảng tỷ lệ thất bại đầy đủ để tránh tạo cảm giác một thuật toán nhanh chỉ vì bỏ cuộc sớm.')
    p('Trong phép đo toàn trận, điểm được so theo thứ tự từ điển (U,D,Q). Nhóm không cộng các thành phần và không gọi một điểm Q cao là tốt hơn khi U hoặc D thấp hơn. Planner tour dùng ngân sách 5 giây/ngày; thời gian lập kế hoạch được đo riêng với mạng và thời gian chờ máy chủ. Practice tự luyện cho phép nộp theo tiến độ người dùng, nên việc nộp thành công không chứng minh hệ thống đã đáp ứng đồng hồ 30 giây ở trận thi tính giờ.')
    p('5.3. Quy trình nộp và lưu nguồn kiểm tra',2)
    p('Trước khi reset một Practice riêng của đội 62, nhóm sao lưu kết quả, trạng thái, kế hoạch và replay hiện có. Sau khi chọn vai trò theo đội hình đã lưu, runner đọc /api/game/day ở đầu từng ngày để nhận vị trí, nhiên liệu và traffics thật. Output được kiểm tra bằng simulator rồi nộp qua /api/game/practice/actions. Điểm cuối trận lấy từ /api/game/result; các bảng điểm ở chương 6 không dùng điểm offline thay cho phản hồi của BTC [9].')
    p('Nhóm chỉ reset và nộp trên lượt Practice riêng của mình. Điểm đội 57 được đọc bằng tài khoản do chủ dự án cung cấp trước đợt đối chiếu; không có thao tác sửa dữ liệu của đội đó. Vì không biết thuật toán, thời gian tính và quá trình chọn lời giải của họ, báo cáo dùng điểm ấy như một mốc tham khảo trên cùng input, không coi là thí nghiệm hai thuật toán cùng điều kiện.')
    p('5.4. Đối chiếu replay và giới hạn bằng chứng',2)
    p('Một frame là trạng thái của cả đội tại một bước thời gian trong replay. Một ngày S step thường có S+1 frame, gồm trạng thái ban đầu. Trong mỗi frame, bộ đối chiếu so số bước, vị trí và nhiên liệu từng xe, số suất, số loại và sự kiện thu hàng ở các trường có trong trace. Ngoài ra, chương trình so kết quả cuối ngày, số frame và trạng thái đường của ngày kế tiếp. Vì vậy, số giá trị được so lớn hơn số frame; chúng không phải các bài kiểm thử độc lập.')
    p('Cùng output đã được BTC nhận được chạy lại trong simulator để đối chiếu, không dùng một lời giải khác có điểm gần giống. Việc không tìm thấy sai khác tăng độ tin cậy trong phạm vi hành trình đã kiểm tra, nhưng không chứng minh simulator đúng trên mọi trạng thái. Mã có hỗ trợ nhiều đội song phần đó chưa được đối chiếu với replay BTC nhiều đội; các kết luận chính hiện chỉ thuộc Practice một đội.')

    p('6. Kết quả và phân tích',1)
    restore(path_results)
    p('6.2. Điểm máy chủ và các cận trên của phiên bản hiện tại',2)
    p('Bảng 11 là kết quả của tour với đội hình được lưu và ngân sách 5 giây/ngày. Mỗi hàng tương ứng một lượt Practice đã hoàn tất 10 ngày trên máy chủ BTC. Cận nguồn cung không cố định vai trò, còn cận theo đội hình giới hạn số suất theo số xe tuần tra; cả hai đều chưa mô tả đầy đủ tính khả thi của lộ trình.')
    p('Bảng 11. Điểm BTC và cận trên trên tám bản đồ Practice',caption=True)
    rows=[]
    for r in data:
        cap=pipeline[r['case']]['new']['bound_for_types']
        # Lấy nguồn cung từ config đã được lưu trong lần nộp thật.
        cfg=json.loads((ROOT/'experiments/tour_api_evidence'/f"{r['case']}_config.json").read_text(encoding='utf8'))
        source=[cap[0],cap[1],sum(s['stocks'] for s in cfg['spots'])*len(cfg['daySteps'])]
        rows.append([r['case'],str(tuple(r['server_score'])),str(tuple(source)),str(tuple(cap))])
    table(['Bài','Điểm BTC (U, D, Q)','Cận theo nguồn cung (U, D, Q)','Cận theo kho và số xe tuần tra (U, D, Q)'],rows,[1.5,4.6,5,4.9])
    p('Cả tám bài đều đạt cận U và D. Điều này xác nhận rằng lời giải thu được mọi loại trong toàn trận và đủ mọi loại ở từng ngày, đồng thời chứng minh hai thành phần đó đã tối ưu trên các input đã thử. Q còn thấp hơn cận nới lỏng không có nghĩa phần chênh lệch chắc chắn thu được: quỹ step, nhiên liệu, vị trí và lịch tiếp tế có thể ngăn việc ghé hết các điểm. Báo cáo chưa xác định Q tối ưu.')
    p('6.3. Đối chiếu với điểm máy chủ của đội khác',2)
    p('Hai bên cùng đạt cận U và D trên các bài đối chiếu, nên thứ tự điểm được quyết định bởi Q. Bảng 12 giữ nguyên ảnh chụp điểm đội 57 đã đọc, còn cột của nhóm là kết quả BTC trong đợt nộp hiện tại. Không lấy điểm offline cao hơn hoặc kết quả của một cấu hình khác để thay vào cột này.')
    p('Bảng 12. So sánh Q khi hai bên đều đạt cận U và D',caption=True)
    table(['Bài','Q của nhóm: BTC','Q đội 57: BTC','Chênh lệch Q','Q/cận sơ bộ theo đội hình'],[
        [r['case'],r['server_score'][2],peers[r['case']]['peer_score'][2],
         f"{r['server_score'][2]-peers[r['case']]['peer_score'][2]:+d}",
         f"{100*r['server_score'][2]/pipeline[r['case']]['new']['bound_for_types'][2]:.1f}%"] for r in data],[1.5,3.1,3.1,2.5,5.8])
    p('Trong lần đối chiếu này, nhóm có bộ điểm cao hơn đội 57 trên cả tám bài. Kết quả thể hiện chất lượng của các lượt nộp cụ thể, không xác định thứ hạng toàn cuộc thi hay chứng minh vượt mọi phương pháp khác. Tỷ lệ ở cột cuối chỉ so Q với giới hạn nới lỏng của bảng 11, không phải tỷ lệ đạt nghiệm tối ưu. Các lần chạy có giới hạn đồng hồ có thể tạo ra tuyến và số suất khác nhau.')
    p('6.4. Thời gian và kết quả đối chiếu replay',2)
    p('Bảng 13. Đội hình, thời gian lập kế hoạch và kiểm chứng máy chủ',caption=True)
    table(['Bài','Vai trò theo thứ tự input (0 tuần tra; 1 tiếp tế)','Thời gian lập kế hoạch lớn nhất/ngày (giây)','Ngày dùng dự phòng','Frame đối chiếu','Sai khác'],[
        [r['case'],str(r['types']),f"{r['max_planning_seconds']:.2f}",r['fallback_days'],r['checked_frames'],r['mismatch_count']] for r in data],[1.5,4.7,3.1,2,2.4,2.3])
    p(f'Tổng cộng, BTC nhận 80 kế hoạch ngày. Bộ đối chiếu đã kiểm tra {fmt(frames)} frame và {fmt(values)} giá trị, không phát hiện sai khác ở các trường được đo. Không ngày nào dùng nhánh Greedy dự phòng. Thời gian lập kế hoạch lớn nhất trong đợt này khoảng {max(r["max_planning_seconds"] for r in data):.2f} giây; số này không bao gồm mạng và thời gian chờ chuyển ngày. Các frame có liên hệ thời gian với nhau, nên không được diễn giải thành {fmt(frames)} input độc lập.')
    p('Kết quả cho thấy lời giải hợp lệ của tour được máy chủ thực thi phù hợp với simulator trên những hành trình đã nộp. Ngân sách 5 giây để lại khoảng trống so với cấu hình 30 giây/ngày, nhưng muốn khẳng định nộp kịp trong trận thi thật vẫn phải kiểm tra thêm đồng hồ, mạng, chọn vai trò và chuyển ngày của chế độ tính giờ.')
    p('6.5. Phân tích một trận đạt đủ loại ở mọi ngày',2)
    q=next(r for r in data if r['case']=='Q01')
    cfg=json.loads((ROOT/'experiments/tour_api_evidence/Q01_config.json').read_text(encoding='utf8'))
    p('Q01 là bản đồ 32×32 với 20 loại udon, bốn xe và mười ngày. Đội hình sử dụng xe đầu làm tiếp tế, ba xe còn lại tuần tra. Bảng 14 trình bày số liệu theo ngày từ replay BTC; các ngày có quỹ step khác nhau, nhưng số loại thu được luôn chạm cận 20. Đây là kết quả của một kế hoạch khả thi, không phải lấy cận lý thuyết thay cho điểm thực tế.')
    p('Bảng 14. Kết quả từng ngày của Q01 từ replay BTC',caption=True)
    table(['Ngày','Step cấp cho mỗi xe','Số loại đã thu trong ngày','Cận số loại/ngày','Số suất thu trong ngày'],[
        [i+1,steps,q['daily_types'][i],q['ud_bound'][0],q['daily_servings'][i]] for i,steps in enumerate(cfg['daySteps'])],[1.5,3.4,4,3.4,3.7])
    pic=doc.add_paragraph();pic.alignment=WD_ALIGN_PARAGRAPH.CENTER;pic.paragraph_format.keep_with_next=True
    pic.add_run().add_picture(str(ROOT/'docs/figures/tour_current_q01.png'),width=Cm(16))
    p('Hình 6. Độ phủ loại và số suất tích lũy của phương pháp hiện tại trên Q01.',caption=True)
    p(f'Ngày 6 chỉ có {cfg["daySteps"][5]} step cho mỗi xe nhưng đội vẫn thu đủ 20 loại. Số suất giữa các ngày khác nhau do quỹ step, vị trí và lịch hẹn không giống nhau, trong khi độ phủ được ưu tiên giữ trước. Tới cuối trận, số suất tích lũy là {q["server_score"][2]}. Việc quan sát từng ngày giúp giải thích ý nghĩa của D = 200: đội không chỉ thu đủ 20 loại một lần rồi dừng, mà duy trì đủ loại ở cả mười ngày.')
    p('Cơ chế chèn chung cho cả đội giúp phân công các điểm cần phủ; lịch hẹn cho phép duy trì nhiên liệu; ước lượng ngày sau giúp giảm nguy cơ dừng ở vị trí bất lợi trước ngày ngắn. Những điều này phù hợp với thiết kế thuật toán và diễn biến replay. Tuy vậy, đợt đo hiện tại chưa tắt riêng từng cơ chế trong cùng ngân sách, nên chưa lượng hóa được phần cải thiện do mỗi thành phần tạo ra.')

    p('7. Đánh giá và kết luận',1)
    p('7.1. Kết quả đạt được',2)
    p('Nhóm đã xây dựng một phương pháp lập tuyến đa xe dựa trên Dijkstra, chèn tham lam, tìm kiếm cục bộ và mô hình tài nguyên. Chương trình nhận input thực tế từng ngày, kiểm tra output rồi nộp tự động. Điểm BTC xác nhận cả tám bài Practice đã thử đạt cận U và D; số suất Q cao hơn mốc đội 57 trong cùng ảnh chụp đối chiếu. Đây là kết quả của thuật toán truyền thống, không sử dụng RL trong luồng chấm chính.')
    p('Phép so sánh năm thuật toán tìm đường cho thấy cần đọc đồng thời chi phí tuyến, mức dùng nhiên liệu, số nhãn và khả năng hoàn thành trong hạn. A* và Dijkstra cho cùng chi phí step trên tập truy vấn chung thành công, trong khi cách mở rộng trạng thái khác nhau. Ở cấp lập kế hoạch, tối ưu từng chặng chỉ là một thành phần; phối hợp điểm ghé, tiếp tế và vị trí cuối ngày mới tạo nên bộ điểm của cả đội.')
    p('7.2. Phạm vi và giới hạn',2)
    p('Tám bản đồ Practice đều là cỡ lớn, chỉ được báo cáo một lượt nộp cho mỗi bài; số mẫu chưa đủ để suy ra độ ổn định trên mọi kích thước, seed hoặc tình huống giao thông. Đội hình được cố định từ thí nghiệm chọn offline. Phần nhiều đội của simulator chưa có replay BTC kiểm chứng, còn Practice tự luyện chưa chứng minh quy trình đáp ứng thời hạn của trận thi thật. Các giới hạn này cần được giữ khi diễn giải kết quả.')
    p('U và D đã đạt cận trên ở các bài được chấm, nhưng Q chưa có nghiệm tối ưu đã biết. Cận Q theo nguồn cung và số xe còn bỏ qua nhiều ràng buộc. Phép đối chiếu replay chỉ xác nhận những giá trị đã kiểm tra trên các kế hoạch đã nộp; nó không thay thế chứng minh tính đúng cho mọi trạng thái của simulator. Điểm của đội tham khảo cũng có thể thay đổi khi họ nộp lại.')
    p('7.3. Hướng phát triển',2)
    p('Nhóm có thể đánh giá nhiều lần chạy với cùng ngân sách, tách tác động của chèn điểm, tìm kiếm cục bộ, phủ loại và ô neo, rồi thử lịch tiếp tế nhiều điểm. Bước tiếp theo cho mục tiêu thi là chạy trọn quy trình trong trận tính giờ thật, bao gồm chọn vai trò và nộp từng ngày. Việc cải thiện Q cần tiếp tục bảo vệ hai tiêu chí loại đã đạt, tránh tăng số suất nhưng làm giảm độ phủ.')
    p('Học tăng cường được giữ như một hướng mở rộng riêng ở phụ lục C. Các số liệu ở đó thuộc mô hình và bộ thực thi của phép thử RL, không phải điểm của phương pháp tour hiện tại. Phần chính tiếp tục tập trung vào mô hình hóa và ứng dụng các thuật toán của môn học.')
    p('7.4. Kết luận',2)
    p('Bài tập lớn đã chuyển bài toán HEXUDON thành một hệ thống có thể nhận dữ liệu, lập kế hoạch và được máy chủ BTC chấm. Giá trị của phương pháp nằm ở cách kết hợp đường đi, phân công điểm ghé, tiếp tế và chuẩn bị ngày tiếp theo, cùng cơ chế kiểm tra thực thi. Kết quả đạt cận U và D trên tám bài đã thử cung cấp bằng chứng cụ thể về hiệu quả ứng dụng, đồng thời báo cáo vẫn phân biệt rõ phần đã xác định tối ưu với phần còn là heuristic.')
    restore(references)
    restore(appendix_a)
    p('Phụ lục B. Chạy chương trình và kiểm tra kết quả hiện tại',1)
    p('B.1. Chạy thuật toán và nộp Practice',2)
    p('Giảng viên dùng Python 3.10 trở lên và mở terminal tại thư mục dự án. Phần thuật toán cùng runner chỉ cần thư viện chuẩn. Token tài khoản được đặt trong biến HEXUDON_TOKEN của tệp .env; không đưa token vào Git. Lệnh dưới dùng tour với ngân sách 5 giây mỗi ngày; nó có chọn vai trò và nộp bài nếu trận đang mở.')
    p('python -X utf8 play_match.py --algorithm tour --mode practice --seconds 5 ID_TRAN',command=True)
    p('Nếu chỉ muốn đọc cấu hình và trạng thái mà không nộp, giảng viên thêm --inspect. Nếu chương trình bị ngắt giữa trận, giảng viên chạy lại cùng lệnh với checkpoint còn giữ. Khi trận kết thúc, script đọc điểm rồi tự dọn dữ liệu tạm. Muốn thử lại trận đã kết thúc, người dùng reset thủ công trên website; mỗi lượt chỉ chạy một tiến trình cho cùng trận.')
    p('Để dùng đúng đội hình của bảng 13 thay vì chọn lại, giảng viên truyền --types sau ID trận. Ví dụ dưới áp dụng đội hình Q01. Không dùng danh sách này cho input khác nếu chưa kiểm tra số xe và vị trí.')
    p('python -X utf8 play_match.py ID_TRAN --algorithm tour --mode practice --seconds 5 --types 1 0 0 0',command=True)
    p('B.2. Đọc điểm máy chủ và đối chiếu replay',2)
    p('Bản tổng hợp tám lượt chấm nằm tại experiments/tour_api_evidence/summary.json, gồm điểm BTC, vai trò, số ngày nhận bài, kết quả từng ngày, thời gian lớn nhất, số frame và số sai khác. Website BTC cho phép đăng nhập tài khoản trong lời mở đầu để kiểm tra thủ công điểm và replay hiện có. Số liệu được lưu theo thời điểm đo; việc nộp lại sau đó có thể thay đổi điểm đang hiển thị.')
    p('Script experiments/verify_tour_server.py tái hiện phép kiểm chứng bằng API: sao lưu lịch sử, reset Practice riêng đội 62, nộp 10 ngày và chạy đối chiếu replay. Nó có thao tác thay đổi lượt Practice, không phải lệnh chỉ xem. Script bỏ qua các bài đã có trong summary.json để tránh ghi đè ngoài ý muốn. Các phản hồi và replay chi tiết được giữ tại máy và bị Git ignore, còn bản tổng hợp không chứa token được đưa cùng mã nguồn [9].')
    p('B.3. Các tệp liên quan đến phương pháp',2)
    p('Bảng B1. Vị trí mã nguồn và dữ liệu của phiên bản được báo cáo',caption=True)
    table(['Tệp hoặc thư mục','Nội dung'],[
        ['src/hexudon/search.py','Năm thuật toán tìm đường có giới hạn tài nguyên'],
        ['src/hexudon/tour.py','Dijkstra theo trọng số, chèn điểm, sửa tuyến, phủ loại và ô neo'],
        ['src/hexudon/team_selection.py','Chọn vai trò bằng mô phỏng các đội hình'],
        ['src/hexudon/model.py; simulator.py','Mô hình dữ liệu, luật thực thi và kiểm tra kế hoạch'],
        ['src/hexudon/live.py; live_api.py; live_storage.py','Điều khiển trận, HTTP và checkpoint'],
        ['experiments/tour_api_evidence/summary.json','Điểm và bằng chứng đối chiếu máy chủ'],
        ['docs/api/openapi.json; docs/auto_submit.md','Đặc tả API và hướng dẫn vận hành'],
        ['tests/','Kiểm thử thuật toán, simulator và runner']],[7.5,8.5])
    p('python -X utf8 -m unittest discover -s tests -v',command=True)
    p('Phần tạo hình có thể cần Matplotlib; phần chỉnh DOCX dùng thư viện tài liệu và Word để cập nhật mục lục, kiểm tra bố cục. Các công cụ này không nằm trong luồng lập kế hoạch khi thi. Không chạy lại các builder báo cáo lịch sử nếu muốn giữ bản DOCX hiện tại, vì chúng có thể tái tạo nội dung của một phạm vi thí nghiệm khác.')
    restore(appendix_c)

    # Chỉ sửa liên kết nội bộ đã lỗi thời; giữ nguyên mọi số liệu RL ở phụ lục riêng.
    for paragraph in doc.paragraphs:
        if 'bảng 18 để quy toàn bộ' in paragraph.text:
            paragraph.text=paragraph.text.replace('điểm Q04 đã tối ưu đội hình ở bảng 18 để quy toàn bộ',
                'điểm của phương pháp tour trong phần chính để quy toàn bộ')
            font(paragraph)
        if paragraph.text.startswith('[9]'):
            paragraph.text='[9] Nhóm thực hiện. Kết quả tour được BTC chấm ngày 20/10/2026: experiments/tour_api_evidence/summary.json; mã kiểm chứng experiments/verify_tour_server.py. Điểm đội 57 được lưu trong experiments/tour_results/server_baseline.json. Giảng viên có thể kiểm tra thủ công điểm và replay của tài khoản phenikaa_2 trên website BTC; dữ liệu hiện trên website có thể thay đổi khi người dùng nộp lại.'
            paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT;font(paragraph,11)
    toc=next(paragraph for paragraph in doc.paragraphs if paragraph.text=='MỤC LỤC')
    repo=find('Link repo Github:')
    node=repo._p.getnext()
    while node is not None and node is not toc._p:
        following=node.getnext()
        if not node.xpath('.//w:t') and not node.xpath('.//w:sectPr'):
            node.getparent().remove(node)
        node=following
    toc.paragraph_format.page_break_before=True
    for root in [doc._element,doc.styles.element]:
        for color in root.xpath('.//w:color'):
            color.set(qn('w:val'),'000000')
            for name in ('themeColor','themeTint','themeShade'):
                color.attrib.pop(qn('w:'+name),None)
    doc.save(path)
    print('Đã thay nội dung kết quả toàn trận bằng phiên bản hiện tại; giữ so sánh tìm đường và phụ lục RL riêng.')


if __name__=='__main__':main()
