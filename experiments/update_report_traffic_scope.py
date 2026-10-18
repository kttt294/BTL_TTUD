"""Ghi rõ phạm vi giao thông một đội trong báo cáo Word hiện có."""

from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
path = ROOT / "output/docx/bao_cao_hexudon.docx"
doc = Document(path)
changes = 0
for paragraph in doc.paragraphs:
    text = paragraph.text
    if text.startswith("Lưu lượng một ô đường được xác định"):
        paragraph.text = text + (
            " Trong phạm vi thực nghiệm hiện tại, chúng em chỉ chạy một đội, nên trạng thái đường "
            "được cập nhật từ lịch sử hiện diện của các xe thuộc chính đội đó trong hai ngày trước; "
            "mẫu số tính lưu lượng bằng 1. Đường vẫn có thể chuyển sang đông đúc hoặc tắc nghẽn, "
            "chứ không được cố định thông thoáng trong mọi ngày. Các bài nộp của tài khoản khác "
            "trong những lượt Practice riêng không tham gia cập nhật giao thông của lượt này.")
        changes += 1
    elif text.startswith("Phần nhiều đội đã được kiểm thử bằng"):
        paragraph.text = (
            "Phần nhiều đội đã được kiểm thử nội bộ bằng các tình huống tính tay về kho độc lập, "
            "tiếp tế trong đội, chia lưu lượng và không sửa đầu vào khi lỗi. Tuy nhiên, các phép đo "
            "so sánh thuật toán và các lượt đối chiếu Practice của BTC trong phạm vi báo cáo đều "
            "chạy một đội. Vì vậy, trạng thái đường ở các kết quả này chỉ phản ánh hoạt động của "
            "xe đội mình. Chúng em chưa đánh giá ảnh hưởng giao thông do đối thủ gây ra và tạm thời "
            "không tổ chức thí nghiệm trận nhiều đội. Khả năng hỗ trợ nhiều đội của mã nguồn được "
            "ghi nhận là một phần mở rộng, chưa được xác minh bằng replay BTC nhiều đội.")
        changes += 1
    elif text.startswith("Ở nhóm toàn trận, chúng em ghi bộ điểm"):
        paragraph.text = text + (
            " Các kết quả này thuộc chế độ một đội: simulator tự tính giao thông từ hoạt động "
            "của xe đội mình, còn khi chạy Practice, chương trình đọc trạng thái giao thông do "
            "BTC trả về cho lượt một đội tương ứng. Chúng em chưa dùng các số liệu này để kết "
            "luận về hiệu quả trong một trận có nhiều đội cùng tác động lên mạng đường.")
        changes += 1
    elif text.startswith("Giới hạn quan trọng nhất là bằng chứng API"):
        paragraph.text = text + (
            " Ngoài ra, toàn bộ đánh giá chiến lược hiện tại giới hạn ở một đội. Báo cáo chưa "
            "đo độ ổn định của thuật toán khi trạng thái đường thay đổi do bài nộp của các đội "
            "khác; thí nghiệm nhiều đội được để ngoài phạm vi của giai đoạn này.")
        changes += 1
assert changes == 4, f"Số đoạn khớp không đúng: {changes}"
doc.save(path)
check = Document(path)
assert len(check.tables) == len(doc.tables)
assert any("mẫu số tính lưu lượng bằng 1" in p.text for p in check.paragraphs)
assert any("tạm thời không tổ chức thí nghiệm trận nhiều đội" in p.text for p in check.paragraphs)
print(f"Updated {changes} paragraphs; document reopened successfully.")
