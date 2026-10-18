"""Tạo báo cáo Word có thể chỉnh sửa từ cùng nội dung của bản PDF."""

from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pdf2image import convert_from_path
from reportlab.graphics import renderPDF

from build_report import make_chart, register_fonts


ROOT = Path(__file__).resolve().parents[1]


def add_text(paragraph, text):
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1])
        else:
            paragraph.add_run(part)


def bookmark(paragraph, name, index):
    begin = OxmlElement("w:bookmarkStart")
    begin.set(qn("w:id"), str(index))
    begin.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(index))
    paragraph._p.insert(0, begin)
    paragraph._p.append(end)


def main():
    report_title = (ROOT / "docs/bao_cao.md").read_text(encoding="utf-8").splitlines()[0][2:]
    output = ROOT / "output/docx"
    output.mkdir(parents=True, exist_ok=True)
    scratch = ROOT / "scratch/docx_review"
    scratch.mkdir(parents=True, exist_ok=True)
    register_fonts()
    renderPDF.drawToFile(make_chart(), str(scratch / "chart.pdf"))
    dependencies = Path(sys.executable).parent.parent
    poppler = dependencies / "native/poppler/Library/bin"
    chart = convert_from_path(str(scratch / "chart.pdf"), dpi=220, poppler_path=str(poppler))[0]
    chart.save(scratch / "chart.png")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    section.header_distance = section.footer_distance = Cm(1)
    section.different_first_page_header_footer = True
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    for text in ("BÁO CÁO BÀI TẬP LỚN", "MÔN THUẬT TOÁN ỨNG DỤNG"):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
        if text.startswith("BÁO"):
            p.paragraph_format.space_before = Pt(75)
    p = doc.add_paragraph(report_title, "Title")
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(24)
    for text in ("Bài toán HEXUDON của Procon Việt Nam 2026",
                 "Thuật toán đồ thị và tìm kiếm có giới hạn tài nguyên", "Tháng 10 năm 2026"):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(30)
    doc.add_page_break()
    doc.add_paragraph("Mục lục", "Title")
    lines = (ROOT / "docs/bao_cao.md").read_text(encoding="utf-8").splitlines()
    chapters = [line[3:] for line in lines if line.startswith("## ")]
    for index, title in enumerate(chapters):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(9)
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), f"chapter_{index}")
        run, properties = OxmlElement("w:r"), OxmlElement("w:rPr")
        font = OxmlElement("w:rFonts")
        font.set(qn("w:ascii"), "Times New Roman")
        font.set(qn("w:hAnsi"), "Times New Roman")
        properties.append(font)
        run.append(properties)
        value = OxmlElement("w:t")
        value.text = title
        run.append(value)
        link.append(run)
        p._p.append(link)
    doc.add_page_break()

    index, chapter_index = 0, 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line or line.startswith(("<!--", "# ", "BÁO CÁO", "Đề tài được xây dựng")):
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(line[3:], "Heading 1")
            bookmark(p, f"chapter_{chapter_index}", chapter_index)
            chapter_index += 1
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("```"):
            while index < len(lines) and not lines[index].startswith("```"):
                p = doc.add_paragraph(lines[index])
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(8.5)
                index += 1
            index += 1
        elif line.startswith("|"):
            rows = [line]
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(lines[index].strip())
                index += 1
            data = [[v.strip() for v in row.strip("|").split("|")]
                    for row in rows if not re.fullmatch(r"[| :\-]+", row)]
            count = len(data[0])
            widths = {8: [25, 65, 55, 65, 55, 100, 55, 55],
                      6: [28, 70, 75, 95, 100, 107], 5: [40, 110, 105, 110, 110],
                      4: [55, 140, 140, 140], 3: [195, 140, 140], 2: [140, 335]}.get(count, [1] * count)
            if count == 5 and data[0][0] == "Bộ replay":
                widths = [160, 40, 70, 100, 105]
            table = doc.add_table(rows=0, cols=count)
            table.autofit = False
            for column, width in zip(table.columns, widths):
                column.width = Cm(16.6 * width / sum(widths))
            borders = OxmlElement("w:tblBorders")
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement("w:" + side)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:" + key), value)
                borders.append(border)
            table._tbl.tblPr.append(borders)
            for row_index, values in enumerate(data):
                row = table.add_row()
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                for col, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                    cell.width = Cm(16.6 * width / sum(widths))
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    p.paragraph_format.line_spacing = 1.1
                    p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(3)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if re.fullmatch(r"[\d./-]+", value) else WD_ALIGN_PARAGRAPH.LEFT
                    add_text(p, value)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.bold = row_index == 0
                    if row_index == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "E6EDF1")
                        cell._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif line.startswith("!["):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(scratch / "chart.png"), width=Cm(16))
            p = doc.add_paragraph("Hình 1 So sánh số phần udon trong cùng ngân sách tính toán")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.italic = True
                run.font.size = Pt(10)
        else:
            text = line
            while index < len(lines) and lines[index].strip() and not lines[index].startswith(("#", "|", "```", "<!--", "![")):
                text += " " + lines[index].strip()
                index += 1
            add_text(doc.add_paragraph(), text)
    doc.core_properties.title = report_title
    doc.core_properties.author = ""
    destination = output / "bao_cao_hexudon.docx"
    doc.save(destination)
    print(destination)


if __name__ == "__main__":
    main()
