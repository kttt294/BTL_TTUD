"""Chèn đối chứng bỏ thành phần vào phần chính, giữ nguyên nội dung phụ lục RL."""

import json
from statistics import mean
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_ablation(doc, root, section, para, table, figure):
    folder=root/'experiments/ablation_results'
    rows=json.loads((folder/'summary.json').read_text(encoding='utf8'))
    audit=json.loads((folder/'audit.json').read_text(encoding='utf8'))
    mechanisms=json.loads((folder/'mechanisms.json').read_text(encoding='utf8'))
    generated=[r for r in rows if r['group']=='generated']
    fmt=lambda x:'('+', '.join(map(str,x))+')'
    body=doc.element.body

    def insert_before(title, writer):
        target=next(p._p for p in doc.paragraphs if p.text==title)
        original=set(body)
        writer()
        created=[node for node in body if node not in original]
        for node in created:
            body.remove(node)
            body.insert(list(body).index(target),node)

    def protocol():
        section('5.5. Đối chứng khi bỏ từng thành phần của bộ lập kế hoạch',2)
        para('Các phép so sánh trước đánh giá cấu hình hoàn chỉnh, nên chưa trả lời rõ '
             'phần phối hợp mục tiêu hoặc cách chọn vai trò đã đóng góp bao nhiêu. Nhóm '
             'bổ sung một thí nghiệm giữ cố định bộ tìm đường A*, luật mô phỏng, số ứng '
             'viên sáu đích, cách xử lý tiếp tế và tiêu chí so điểm. Mỗi biến thể chỉ '
             'thay một quyết định được nêu trong bảng 24 ở mục 6.7. Các mức này là các nhánh '
             'đối chứng quanh cùng một bộ điều khiển, không phải một chuỗi cải tiến '
             'mà mọi mức sau đều phải tốt hơn mức trước.')
        para('G0 vẫn biết toàn bộ trạng thái thật ở đầu ngày nhưng không biết tuyến '
             'các xe khác đang được lập: trước khi xử lý một xe, chương trình đặt lại '
             'kho dự kiến, tập loại dự kiến trong ngày và tập loại dự kiến toàn trận '
             'về thông tin đầu ngày. G1 giữ các thông tin này sau từng tuyến, nhờ đó '
             'xe sau có thể tránh điểm đã được xe trước đặt chỗ. Cả hai vẫn được chấm '
             'bằng thực thi đồng thời và dùng chung tồn kho thật. Phép đặt chỗ theo '
             'thứ tự xe chưa mô hình hóa chính xác mọi thời điểm đến, nên vẫn có thể '
             'xuất hiện lượt ghé điểm đã hết hàng dù G1 đã bật phối hợp.')
        para('S* xét toàn tuần tra và bốn phương án có đúng một xe tiếp tế, giữ nguyên '
             'vị trí xuất phát. Chương trình mô phỏng hoàn tất cả năm phương án bằng '
             'cùng Greedy có phối hợp rồi chọn theo (U,D,Q). Đây là tập vai trò giới '
             'hạn, không phải vét cạn cả 15 cách có ít nhất một xe tuần tra, và không '
             'phải nghiệm tối ưu của bài toán đầy đủ. Chi phí của S* bao gồm năm lần '
             'đánh giá toàn trận; điểm của nó không được so như thể chỉ tốn một lần '
             'lập kế hoạch tương đương các cấu hình cố định.')
        para('Tập đo có chín bản đồ sinh, gồm ba seed 42000, 42001 và 42002 ở mỗi '
             'kích thước 8×8, 16×16 và 32×32. Mỗi trận có bốn ngày, 80 step/ngày, '
             'bình nhiên liệu 40 và bốn vị trí xuất phát. Nhóm bổ sung hai input minh '
             'họa 8×8 để quan sát trường hợp các cấu hình bằng điểm và trường hợp '
             'nhiên liệu rất thấp; hai input này không gộp vào tỷ lệ thắng trên tập '
             'bản đồ sinh. Toàn bộ điểm của lượt đo mới là điểm offline một đội, '
             'không phải điểm máy chủ BTC hoặc thử nghiệm nhiều đội.')
        para('Để giảm nhiễu do deadline, bản sao planner dùng trong phép đo không '
             'ngắt tìm kiếm theo đồng hồ. Nó giữ nguyên quy tắc chọn đích và A* của '
             'planner chính; hai kiểm thử mới đối chiếu bản sao với bản gốc khi đủ '
             'thời gian và kiểm tra một trường hợp tranh kho. Quyết định trên cùng '
             'input vì vậy có thể tái hiện, còn số giây vẫn phụ thuộc máy. Phép đo '
             'này phục vụ phân tích thành phần, không chứng minh mọi cấu hình sẽ '
             'đáp ứng hạn 30 giây. Đánh giá dưới deadline thực vẫn dùng các bảng '
             'đã trình bày ở mục 6.6.')

    def results():
        section('6.7. Đóng góp riêng của phối hợp, nhìn trước và chọn vai trò',2)
        table('Bảng 24. Thành phần được thay đổi và đối chứng tương ứng',
              ['Mã biến thể', 'Cách chọn mục tiêu và vai trò', 'Biến thể dùng để đối chiếu',
               'Thành phần được kiểm tra'],
              [['G0','Greedy, bốn xe tuần tra lập tuyến độc lập','G1','Tác dụng của việc phối hợp giữa các xe'],
               ['G1','Greedy có phối hợp, bốn xe tuần tra','G0','Chia sẻ loại đã dự kiến thu và đặt chỗ tồn kho'],
               ['L1','Lookahead hai điểm có phối hợp, bốn xe tuần tra','G1','Lợi ích và giới hạn của việc nhìn trước'],
               ['S1','Greedy có phối hợp, cố định xe cuối làm tiếp tế','G1','Đánh đổi giữa xe thu hàng và nguồn tiếp nhiên liệu'],
               ['S*','Greedy có phối hợp, chọn tốt nhất trong năm vai trò','S1','Lợi ích của lựa chọn xe tiếp tế theo vị trí xuất phát']],
              [1.5,6,2.8,5.7],10.5)
        labels={'coordinated':'Phối hợp G1 so với độc lập G0',
                'lookahead':'Nhìn trước L1 so với Greedy G1',
                'last_supply':'Xe cuối tiếp tế S1 so với toàn tuần tra G1',
                'selected':'Chọn vai trò S* so với xe cuối tiếp tế S1'}
        aggregate=[]
        for method in labels:
            entries=[r for r in audit['comparisons'] if r['method']==method]
            new_searches=mean(r['mean_searches_new'] for r in entries)
            if method=='selected':
                new_searches=mean(sum(t['searches'] for t in json.loads(
                    (folder/f"{r['case']}_detail.json").read_text(encoding='utf8'))['selection_trials'].values())
                    for r in generated)
            aggregate.append([labels[method],sum(r['n'] for r in entries),
                *[sum(r[k] for r in entries) for k in ('wins','ties','losses')],
                f"{mean(r['mean_searches_baseline'] for r in entries):.0f} / "
                f"{new_searches:.0f}"])
        table('Bảng 25. Kết quả tách thành phần trên chín bản đồ sinh',
              ['Thay đổi được đánh giá', 'Số cặp input', 'Tốt hơn đối chứng',
               'Bằng điểm đối chứng', 'Kém hơn đối chứng',
               'Truy vấn A* trung bình cả trận: đối chứng / biến thể'],
              aggregate,[4.8,1.8,2.1,2.1,2.1,3.1],10.5)
        table('Bảng 26. Điểm từng input khi giữ cố định bộ tìm đường và luật mô phỏng',
              ['Kích thước / seed', 'G0: độc lập (U,D,Q)', 'G1: phối hợp (U,D,Q)',
               'L1: nhìn trước (U,D,Q)', 'S1: xe cuối tiếp tế (U,D,Q)', 'S*: chọn vai trò (U,D,Q)'],
              [[f"{r['size']}×{r['size']} / {r['case'].split('_')[-1]}",
                *[fmt(r['scores'][k]) for k in ('independent','coordinated','lookahead','last_supply','selected')]]
               for r in generated],[2.1,2.75,2.75,2.75,2.75,2.9],10.5)
        para('Cột truy vấn của S* tính tổng công việc của cả năm phương án '
             'vai trò được xét; các biến thể cố định chỉ tính một trận. '
             'Vì vậy, lợi ích điểm của S* đi kèm chi phí tìm kiếm trước '
             'trận lớn hơn, không phải một cải thiện miễn phí với cùng '
             'ngân sách tính toán.')
        para('G1 tốt hơn G0 ở tám bản đồ và kém ở một bản đồ. Điều này cho thấy '
             'phối hợp giữa các xe là một đóng góp thực nghiệm rõ trong bộ dữ liệu '
             'nhỏ này, nhưng quy tắc đặt chỗ hiện tại chưa bảo đảm tăng điểm. Khi '
             'loại một mục tiêu vì xe trước đã chọn, xe sau có thể chuyển sang '
             'một tuyến khác và kết thúc ngày ở vị trí bất lợi cho ngày tiếp theo.')
        para('L1 thắng G1 ở ba input, bằng ở một và thua ở năm, mặc dù số truy vấn '
             'A* trung bình tăng từ khoảng 322 lên 731 mỗi trận. Do phép đo đã bỏ '
             'deadline, những trường hợp thua này không thể chỉ được giải thích '
             'bằng việc Lookahead bị hết thời gian sớm. Hàm đánh giá cặp điểm vẫn '
             'chỉ nhìn trước một đoạn ngắn, trong khi vị trí kết thúc và lượng '
             'nhiên liệu còn lại ảnh hưởng tới toàn bộ các ngày sau.')
        para('S1 thắng G1 ở tám input nhưng thua ở một input 32×32. Xe tiếp tế '
             'giúp kéo dài hoạt động nhưng cũng làm mất một xe thu hàng; nếu xe '
             'được đổi vai trò đang ở vùng có loại udon khó thay thế, điểm U có '
             'thể giảm dù D và Q tăng. Khi xét đủ năm phương án giới hạn, S* '
             'thắng S1 ở năm input và bằng ở bốn. Đây là tác dụng của việc thử '
             'vai trò theo vị trí xuất phát với cùng bộ điều khiển, không phải '
             'bằng chứng rằng đội hình đã tối ưu toàn cục.')
        para('Trên chín bản đồ sinh, tổng thời gian đánh giá năm phương án của S* '
             f"nằm từ {min(r['selection_seconds'] for r in generated):.2f} đến "
             f"{max(r['selection_seconds'] for r in generated):.2f} giây trên máy đo. "
             'Đây là chi phí mô phỏng chọn vai trò trước trận, chưa phải ngân '
             'sách lập kế hoạch cho một ngày. S* không kém S1 trong phép đo này '
             'vì chính phương án S1 nằm trong tập được xét và mọi phương án được '
             'mô phỏng hoàn tất, xác định, trên cùng môi trường. Kết luận đó '
             'không tự mở rộng sang giao thông do các đội khác tạo ra.')
        para(f"Đợt đo chạy {audit['unique_games']} trận khác nhau với "
             f"{audit['verified_days']} kế hoạch ngày. Các cấu hình dùng chung vai "
             'trò được tái sử dụng khi tổng hợp thay vì chạy lại và đếm thêm mẫu. '
             'Mọi kế hoạch đều hợp lệ, không dùng dự phòng và chấm lại khớp điểm. '
             'Mỗi input chỉ có một lượt đo xác định; chín bản đồ chưa đủ để '
             'khẳng định thứ hạng chung trên mọi dạng đề.')
        para('Input minh họa shared_stock cho cả năm biến thể cùng điểm (4,8,8). '
             'Bản đồ này có bốn điểm thuộc bốn loại, mỗi điểm một suất mỗi ngày '
             'và hai ngày thi, nên nguồn cung chỉ cho phép tối đa bốn loại '
             'toàn trận, tám loại qua các ngày và tám suất. Các biến thể đều '
             'đã thu hết nguồn cung trong cấu hình này; việc phối hợp hoặc '
             'đổi vai trò không còn dư địa tăng điểm. Trường hợp bằng nhau '
             'này được giữ lại cùng trường hợp thiếu nhiên liệu, thay vì '
             'chỉ chọn ví dụ có cải thiện.')
        section('6.8. Phân tích nguyên nhân từ hành trình và trạng thái cuối ngày',2)
        para('Trường hợp phối hợp mang lại lợi ích.',bold=True)
        para('Ở bản đồ 8×8, seed 42000, cả G0 và G1 đều thu đủ tám loại trong ngày '
             'đầu. Tuy nhiên, G0 có 27 lượt một xe lần đầu trong ngày ghé điểm '
             'đã hết hàng, còn G1 có bốn lượt. Chương trình đếm chỉ số này từ '
             'replay, theo đúng thứ tự xử lý xe và tồn kho thật, không suy ra '
             'từ các đích đã đặt chỗ. Cuối ngày thứ hai, cả bốn xe của G0 hết '
             'nhiên liệu; hai xe của G1 vẫn còn lần lượt ba và sáu đơn vị. '
             'Số loại thu theo ngày của G0 là [8,5,2,2], còn G1 là [8,8,7,4], '
             'nên điểm cuối tăng từ (8,17,49) lên (8,27,67).')
        para('Phép đối chứng chỉ thay cơ chế chia sẻ mục tiêu nên có thể quy '
             'chênh lệch cho cơ chế đó trong cấu hình đang xét. Tuy vậy, nhóm '
             'không coi riêng việc giảm lượt ghé kho rỗng là lời giải thích '
             'đầy đủ: tuyến thay đổi còn làm thay đổi nhiên liệu, vị trí cuối '
             'ngày và giao thông của chính đội. Các số đo này giúp nối quyết '
             'định của thuật toán với kết quả, thay vì chỉ kết luận rằng phối '
             'hợp luôn tốt hơn.')
        table('Bảng 27. Phối hợp giảm ghé kho rỗng nhưng chưa bảo đảm điểm tốt hơn',
              ['Bản đồ được phân tích', 'Biến thể', 'Điểm cuối trận (U,D,Q)',
               'Lượt ghé điểm đã hết hàng', 'Số loại thu lần lượt qua bốn ngày'],
              [[f'{size}×{size}, seed 42000',code,
                fmt(mechanisms[f'generated_{size}_42000'][method][-1]['score']),
                sum(r['empty_stock_visits'] for r in mechanisms[f'generated_{size}_42000'][method]),
                str([r['daily_types'] for r in mechanisms[f'generated_{size}_42000'][method]])]
               for size in (8,16) for method,code in [('independent','G0'),('coordinated','G1')]],
              [3.3,1.6,3.3,3.5,4.3],10.5)
        para('Phản ví dụ cho việc chỉ tối ưu đặt chỗ.',bold=True)
        para('Ở bản đồ 16×16 cùng seed 42000, G1 loại được cả ba lượt ghé kho '
             'rỗng của G0, nhưng điểm lại đổi từ (12,25,51) thành (12,24,55). '
             'Hai cấu hình bằng U, nên D thấp hơn khiến G1 thua dù Q cao hơn. '
             'Sau ngày đầu, G0 còn nhiên liệu [0,0,7,1], còn G1 là [0,0,4,0]. '
             'Ngày thứ hai, G0 thu thêm năm loại còn G1 chỉ thu bốn; các ngày '
             'sau đều thu bốn loại. Kết quả cho thấy giữ chỗ theo thứ tự xe có '
             'thể tránh tranh kho nhưng chưa cân đối vị trí và tài nguyên cho '
             'cả đội. Muốn cải tiến tiếp, nhóm cần xét chi phí cơ hội của việc '
             'giao một đích cho một xe, thay vì chỉ loại đích đó khỏi xe sau.')
        figure('ablation_daily_types.png',
               'Hình 8. Diễn biến D khi bật phối hợp trên hai input có kết quả trái chiều.',16)
        para('Lookahead có thể chọn tốt trước mắt nhưng kết thúc ở vị trí bất lợi.',bold=True)
        para('Input minh họa scarce_fuel có bản đồ 8×8 toàn đồng bằng, bốn xe '
             'tuần tra, bốn ngày với 32 step/ngày và bình nhiên liệu tám đơn '
             'vị. Sáu điểm bán có vị trí [19,20,27,28,35,36], loại tương ứng '
             '[0,1,2,0,1,2] và kho bốn suất mỗi điểm mỗi ngày. Các xe xuất '
             'phát tại [9,10,17,18]. Input này được thiết kế để làm rõ trạng '
             'thái sau khi cạn nhiên liệu, không phải một mẫu ngẫu nhiên đại '
             'diện cho đề BTC.')
        positions=mechanisms['scarce_fuel']
        cfg=json.loads((folder/'scarce_fuel_input.json').read_text(encoding='utf8'))
        brands={s['pos']:s['brand'] for s in cfg['spots']}
        g=positions['coordinated'][1]['final_positions']
        l=positions['lookahead'][1]['final_positions']
        table('Bảng 28. Vị trí sau ngày thứ hai khi cả bốn xe đều đã hết nhiên liệu',
              ['Chỉ số xe theo input', 'Ô dừng của Greedy G1', 'Loại tại ô dừng G1',
               'Ô dừng của Lookahead L1', 'Loại tại ô dừng L1'],
              [[car,g[car],brands.get(g[car],'Không có'),l[car],brands.get(l[car],'Không có')]
               for car in range(4)],[2.8,3.3,3.3,3.3,3.3],10.5)
        para('Sau hai ngày, tất cả xe của cả hai cấu hình đều hết nhiên liệu. '
             'G1 dừng tại các điểm thuộc đủ ba loại, trong khi L1 dừng tại '
             'các điểm chỉ thuộc loại 0. Do kho được bổ sung mỗi ngày và xe '
             'đang đứng tại điểm bán vẫn thu được hàng, G1 tiếp tục thu ba '
             'loại/ngày còn L1 chỉ thu một loại/ngày. Hai cấu hình cùng thu '
             '40 suất trong trận nhưng đạt lần lượt (3,12,40) và (3,8,40). '
             'Đây là một cơ chế cụ thể khiến nhìn trước hai điểm vẫn thua: '
             'hàm đánh giá chưa tính đủ giá trị của vị trí dừng khi tài '
             'nguyên không còn cho phép di chuyển ở những ngày cuối.')
        para('Khi đổi xe cuối thành tiếp tế trên chính input này, S1 đạt '
             '(3,12,72). Như vậy, vai trò tiếp tế có thể tăng số suất ngay '
             'cả khi số loại đã bằng nhau. Tuy nhiên, trên input 32×32 seed '
             '42000, G1 đạt (12,21,33) còn S1 đạt (11,41,73): tăng mạnh D và '
             'Q vẫn không bù được việc mất một loại. Replay cho thấy S1 bỏ '
             'lỡ loại 6 trong cả trận. Khi S* đổi xe chỉ số 1 thành tiếp '
             'tế thay vì xe chỉ số 3, điểm đạt (12,42,69). Ví dụ này cho '
             'thấy quyết định chọn xe nào làm tiếp tế quan trọng hơn việc '
             'chỉ quy định đội hình phải có một xe tiếp tế.')
        para('Các ví dụ trên được chọn từ dữ liệu đã đo để giải thích cơ chế, '
             'không dùng riêng chúng để suy ra tỷ lệ thắng tổng quát. Nhận '
             'xét tổng hợp vẫn dựa trên chín input ở bảng 25 và bảng 26. '
             'Nhóm lưu input, chuỗi lệnh và trạng thái từng ngày để người '
             'đọc có thể kiểm tra cả ví dụ tăng điểm lẫn phản ví dụ giảm điểm.')

    def reproduction():
        section('B.4. Tái hiện phép đo tách thành phần và kiểm tra ví dụ',2)
        for command in ('python -X utf8 experiments/benchmark_ablation.py',
                        'python -X utf8 experiments/audit_ablation.py',
                        'python -X utf8 experiments/analyze_ablation.py',
                        'python -X utf8 experiments/plot_ablation.py'):
            para(command).alignment=WD_ALIGN_PARAGRAPH.LEFT
        para('Ba lệnh đầu chỉ dùng thư viện chuẩn; lệnh cuối cần Matplotlib để '
             'vẽ hình. Thư mục experiments/ablation_results giữ 11 input cùng '
             'chi tiết bảy trận khác nhau trên mỗi input, bảng tổng hợp và '
             'nhật ký phân tích. Lệnh audit chấm lại 77 trận và 294 kế hoạch '
             'ngày; lệnh analyze kiểm tra các sự kiện thu hàng trước khi '
             'đếm lượt ghé kho rỗng. Không lệnh nào gọi API hoặc đặt lại '
             'Practice của đội. Chạy benchmark lần nữa sẽ ghi đè dữ liệu '
             'trong thư mục; quyết định không phụ thuộc deadline nhưng số '
             'giây vẫn có thể đổi theo máy.')

    insert_before('6. Kết quả và phân tích thuật toán',protocol)
    insert_before('7. Đánh giá giới hạn và kết luận',results)
    insert_before('Phụ lục C. Thử nghiệm mở rộng với học tăng cường',reproduction)
    # Cập nhật hướng tiếp theo vì đối chứng cùng bộ điều khiển nay đã được thực hiện.
    for p in doc.paragraphs:
        if p.text.startswith('Sau phép đo nhiều ngân sách ở mục 6.6, hướng tiếp theo'):
            p.text=('Các phép đo ở mục 6.7 và 6.8 cho thấy phối hợp và lựa chọn vai trò '
                    'có đóng góp nhưng vẫn có phản ví dụ. Hướng tiếp theo là cải thiện '
                    'cách giao mục tiêu theo chi phí cơ hội và giá trị vị trí cuối ngày, '
                    'sau đó đo lại dưới deadline trên nhiều seed. Nhóm cũng có thể thử '
                    'tìm kiếm cục bộ có kiểm tra tài nguyên; báo cáo chưa cài 2-opt, '
                    '3-opt hoặc beam search. Mỗi cải tiến cần có đối chứng riêng thay '
                    'vì chỉ bổ sung tên thuật toán.')
    return len(doc.tables)
