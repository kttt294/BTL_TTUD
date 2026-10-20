"""Cập nhật có chọn lọc DOCX hiện có bằng kết quả tour đã được BTC chấm."""

from copy import deepcopy
import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def update_report():
    path = ROOT/'output/docx/bao_cao_hexudon.docx'
    doc = Document(path)
    data = json.loads((ROOT/'experiments/tour_api_evidence/summary.json').read_text(encoding='utf8'))
    baseline = json.loads((ROOT/'experiments/tour_results/server_baseline.json').read_text(encoding='utf8'))
    pipeline = json.loads((ROOT/'experiments/tour_results/pipeline.json').read_text(encoding='utf8'))
    if len(data) != 8 or any(r['accepted_days'] != 10 or r['mismatch_count'] for r in data):
        raise ValueError('Cần đủ 8 trận đã được BTC chấm và replay không sai khác trước khi viết kết quả.')
    if any(p.text.startswith('4.8. Lập tuyến cả đội') for p in doc.paragraphs):
        raise ValueError('Báo cáo đã có mục tour; không chèn lặp.')
    headings = {p.text: p for p in doc.paragraphs if p.style.name.startswith('Heading')}
    body = next(p for p in doc.paragraphs if p.text.startswith('HEXUDON là bài toán'))
    heading = headings['4.7. Quy hoạch động làm đối chứng tối ưu trên bài nhỏ']
    def find(prefix):
        return next(p for p in doc.paragraphs if p.text.startswith(prefix) and not p.style.name.startswith('toc'))
    def font(p, size=12.5, bold=False):
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0,0,0)
            run.font.bold = bold
    def paragraph(anchor, text, title=False, caption=False):
        template = heading if title else body
        p = anchor.insert_paragraph_before(text, style=template.style)
        if template._p.pPr is not None:
            p._p.insert(0, deepcopy(template._p.pPr))
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.keep_with_next = title or caption
        p.paragraph_format.keep_together = True
        if text.startswith('6.9.'):
            p.paragraph_format.page_break_before = True
        if text.startswith('[9]'):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if caption:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        font(p, size=11 if text.startswith('[9]') else 12.5, bold=title or caption)
        return p
    def replace(prefix, text):
        p = find(prefix)
        p.text = text
        font(p)
    def table(anchor, headers, rows, widths):
        available=(doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin)/360000
        widths=[w*available/sum(widths) for w in widths]
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        t.autofit = False
        for column,width in zip(t.columns,widths):
            column.width=Cm(width)
        for cells, values in [(t.rows[0].cells,headers)]:
            for cell, value in zip(cells,values): cell.text = str(value)
        for values in rows:
            for cell,value in zip(t.add_row().cells,values): cell.text=str(value)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i,row in enumerate(t.rows):
            cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.1
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.keep_with_next = i < len(t.rows)-1
                    font(p,10.5,i==0)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEDED');cell._tc.get_or_add_tcPr().append(shade)
        anchor._p.addprevious(t._tbl)
        return t

    # Phần mở đầu và phạm vi phản ánh phương pháp tốt nhất, giữ các phép đo cũ làm đối chứng.
    replace('Ở lớp tìm đường, nhóm đối chiếu',
        'Ở lớp tìm đường, nhóm đối chiếu DFS, BFS, Greedy Best-First, Dijkstra và A* trên trạng thái có tài nguyên. Ở lớp lập kế hoạch, nhóm giữ Greedy và Lookahead hai điểm làm đối chứng, sau đó xây dựng phương pháp lập tuyến cả đội bằng chèn tham lam, tìm kiếm cục bộ phá và dựng lại, kết hợp Dijkstra và lịch hẹn tiếp tế. Phương pháp mới còn chuẩn bị vị trí cuối ngày để bảo vệ khả năng thu đủ loại udon trong những ngày có ít step. Các phép đo cũ vẫn được giữ: ở ngân sách 0,25 giây/ngày, Lookahead thắng 3 và thua 27 trong 30 cặp; kết quả này chỉ áp dụng cho hai đối chứng ban đầu.')
    frames=sum(r['checked_frames'] for r in data)
    values=sum(r['checked_values'] for r in data)
    ud=sum(r['server_score'][:2]==r['ud_bound'] for r in data)
    wins=sum(tuple(r['server_score'])>tuple(baseline[r['case']]['peer_score']) for r in data)
    replace('Hai cấu hình phân vai cho các xe',
        f'Ở đợt kiểm chứng mới, nhóm nộp phương pháp lập tuyến cả đội lên website BTC trên 8 bản đồ Practice khác nhau, với ngân sách 5 giây/ngày và đội hình lấy từ lượt chọn offline. Máy chủ đã nhận đủ 80 kế hoạch ngày; phương pháp đạt cận trên của cả số loại toàn trận U và tổng số loại theo ngày D ở {ud}/8 bài. Nhóm đối chiếu thêm {frames} khung trạng thái và {values} giá trị trong replay; không phát hiện sai khác ở các trường được kiểm tra. Bộ dữ liệu Practice khảo sát có kích thước 30×30 hoặc 32×32, nên nhóm vẫn dùng bản đồ tự sinh để đánh giá thêm các kích thước khác trong phạm vi đề bài từ 8 đến 32.')
    replace('Ngoài ra, các cận trên',
        'Các cận trên được tính riêng cho từng thành phần điểm. Khi lời giải khả thi do BTC chấm đạt cận U và D, nhóm có thể kết luận hai thành phần đó đã tối ưu trên bài tương ứng; kết luận này không mở rộng thành tối ưu toàn bộ bộ điểm vì số suất Q vẫn chưa có chứng minh tối ưu. Nhóm cũng dùng quy hoạch động tìm nghiệm tối ưu chính xác trên 24 bài thu nhỏ 2×2 và 3×3 một xe, đồng thời giữ 240 lượt đo độ nhạy ngân sách trên 30 bản đồ tự sinh để phân tích các đối chứng ban đầu.')
    replace('Các câu hỏi chính được đặt ra là:',
        'Các câu hỏi chính là: A* giảm lượng tìm kiếm so với Dijkstra ra sao; Lookahead có tăng điểm so với Greedy trong cùng ngân sách không; lựa chọn vai trò và phối hợp tiếp tế ảnh hưởng thế nào; và việc lập tuyến cả đội, cải thiện cục bộ cùng chuẩn bị ngày kế tiếp có giúp đạt cận độ phủ trên máy chủ BTC hay không. Quy hoạch động trên bài thu nhỏ cung cấp thêm mốc nghiệm tối ưu. Mỗi câu hỏi được đánh giá ở đúng lớp quyết định, tránh lấy đường đi ngắn để suy ra điểm toàn trận tối ưu.')
    replace('Sản phẩm gồm năm thuật toán',
        'Sản phẩm gồm năm thuật toán tìm đường có xét tài nguyên, Greedy và Lookahead làm đối chứng, phương pháp lập tuyến cả đội có cải thiện cục bộ, bộ chọn đội hình, bộ tính cận tiếp cận, simulator và validator. Nhóm cũng cung cấp công cụ nộp tự động, lưu số liệu và đối chiếu từng khung replay với máy chủ BTC.')
    paragraph(find('4.1. Greedy'), 'Các mục 4.1–4.7 trình bày các đối chứng và công cụ nền tảng. Mục 4.8 bổ sung phương pháp lập tuyến cả đội của phiên bản mới; các thành phần này được đánh giá trên máy chủ ở mục 6.9.')
    paragraph(find('6.3. Kết quả chọn đội hình'), 'Các bảng trong mục 6.3 ghi nhận bộ chọn đội hình ban đầu đi cùng Greedy/Lookahead, không phải kết quả tốt nhất của phiên bản tour mới. Nhóm giữ chúng để thể hiện đối chứng và giới hạn của cách chọn vai trò; bảng điểm BTC của phiên bản mới nằm ở mục 6.9.')

    a=find('5. Thiết kế thực nghiệm')
    paragraph(a,'4.8. Lập tuyến cả đội bằng chèn tham lam và tìm kiếm cục bộ',True)
    for text in [
        'Hạn chế quan trọng của Greedy ban đầu là mỗi xe quyết định điểm kế tiếp trên một trạng thái hiện tại, trong khi việc tiêu nhiên liệu và vị trí dừng có thể làm giảm số loại thu được ở những ngày sau. Phương pháp mới, được đặt tên tour trong mã nguồn, biểu diễn mỗi xe tuần tra bằng một dãy điểm ghé và đánh giá đồng thời tập tuyến của cả đội. Phương pháp vẫn thuộc nhóm heuristic của thuật toán ứng dụng, không dùng mô hình học máy.',
        'Ở lớp đường đi, chương trình dùng Dijkstra để tạo cây đường ngắn nhất theo trọng số a×step + b×nhiên liệu. Ba cặp trọng số (1; 0,001), (1; 1) và (0,001; 1) lần lượt thiên về đi nhanh, cân bằng và tiết kiệm nhiên liệu. Cây được tính khi cần và lưu lại trong một ngày. Khác với A* có nhãn tài nguyên của đối chứng, mỗi trọng số chỉ cung cấp một phương án đường đại diện; đây không phải liệt kê đầy đủ mọi đường Pareto về step và nhiên liệu. Khi ghép các chặng, chương trình vẫn kiểm tra tài nguyên và tính cả các điểm udon nằm trên đường.',
        'Để dựng lời giải ban đầu, chương trình xét cách chèn một điểm vào tuyến của các xe tuần tra. Mức ưu tiên lần lượt là loại mới toàn trận, loại mới trong ngày, thêm suất còn trong kho và mức tăng chi phí nhỏ hơn. Các lượt chèn cập nhật chung số xe đã được phân công ghé mỗi điểm. Với tập xe đã ghé điểm p, số suất dự kiến tại điểm đó là min(kho tại p, số xe ghé p); mỗi xe chỉ được nhận tại điểm đó một lần trong ngày. Cách đánh giá chung giúp tránh việc một xe tối ưu tuyến riêng nhưng làm mất cơ hội của cả đội.',
        'Có thể hình dung thao tác chèn bằng đoạn tuyến a → b. Khi muốn ghé thêm điểm p, chương trình thử thay đoạn đó bằng a → p → b, với phần chi phí tăng gần bằng c(a,p) + c(p,b) − c(a,b). Các vị trí chèn được sàng lọc để tiết kiệm thời gian, rồi tuyến mới được chạy thử lại để kiểm tra quỹ step, nhiên liệu, tiếp tế và các điểm đi qua. Sau nhiều lượt chèn, lời giải là các dãy điểm cho cả đội chứ không còn một quyết định điểm đích độc lập. Do chỉ khảo sát một phần vị trí và kiểu đường, đây vẫn là heuristic, không phải bộ giải chính xác bài toán định tuyến xe.',
        'Sau bước dựng, chương trình dùng tìm kiếm cục bộ kiểu phá và dựng lại. Mỗi lần thử, nó bỏ ngẫu nhiên một vài điểm hoặc một đoạn tuyến, đôi khi đổi kiểu đường của một xe, rồi chèn lại với nhiễu nhỏ ở chi phí. Phương án mới chỉ được nhận khi bộ giá trị đánh giá không giảm. Khi β = 0 như cấu hình thực nghiệm, bộ giá trị gồm loại mới toàn trận, số loại trong ngày, số suất, âm nhiên liệu đã tiêu và âm step. Hai thành phần cuối chỉ phá hòa và không lấn át ưu tiên điểm của đề.',
        'Đội hình không có xe tiếp tế được chia mức nhiên liệu sử dụng theo tỷ lệ quỹ step hôm nay trên tổng quỹ step còn lại; đây là heuristic tránh tiêu cạn bình ở đầu trận. Với đội có tiếp tế, chương trình thử một số điểm hẹn thay vì cố định một điểm trung tâm cho cả trận. Xe tuần tra có thể rẽ vào điểm hẹn để nạp nhiên liệu khi chặng tiếp theo và đường quay lại không còn đủ nhiên liệu. Việc đổi số xe tiếp tế vẫn là đánh đổi giữa số xe lấy hàng và khả năng hoạt động lâu dài.',
        'Để bảo vệ D trong ngày ngắn, hàm cover_types dựng lại nhiều khung tuyến chỉ nhắm phủ loại, sau đó mới chèn thêm suất vào khung tốt. Khi còn ngày sau, hàm find_anchors ước lượng khả năng phủ loại từ vị trí và nhiên liệu cuối ngày, rồi tìm các ô neo để xe dừng thuận lợi hơn. Chương trình có thể nhận một tuyến giảm số suất hôm nay nếu độ ưu tiên loại hôm nay và độ phủ dự đoán ngày mai tốt hơn. Giao thông ngày mai được tạm lấy bằng giao thông hôm nay trong phép ước lượng này; đây là dự báo heuristic, không phải dữ liệu tương lai từ máy chủ.',
        'Bộ chọn vai trò cho tour chạy các ứng viên bằng nhiều tiến trình, sàng lọc rồi đánh giá lại tối đa bốn đội hình đứng đầu khi còn thời gian. Nếu không có lượt mô phỏng trọn trận, nó có thể chọn theo số ngày đã chạy và điểm hiện có; trường hợp này phải được hiểu là ước lượng chưa hoàn chỉnh. Trong đợt nộp ở mục 6.9, nhóm cố định đội hình đã được lưu từ thí nghiệm offline, để tách việc chấm tuyến khỏi biến động của bước chọn vai trò.',
        'Giả sử một ngày cần cây từ R ô nguồn, đồ thị có V ô và E cạnh, ba trọng số Dijkstra có chi phí gần O(3R(V+E)log V) cho các cây thực sự được tạo. Phần chèn và sửa tuyến còn phụ thuộc số xe, số điểm, số vị trí chèn được xét và số vòng cải thiện, nên không thể quy toàn bộ phương pháp về một lần tìm đường. Tìm kiếm dừng theo ngân sách, nhưng lượt dựng đầu và một số phép ước lượng không bị cắt ở mọi thao tác; ngân sách là giới hạn mềm. Validator kiểm tra lời giải trước khi nộp; nếu tour lỗi, hệ thống có nhánh dự phòng Greedy và số lần dùng nhánh này được thống kê riêng.'
    ]:paragraph(a,text)

    a=find('6. Kết quả và phân tích')
    paragraph(a,'5.6. Quy trình đánh giá phương pháp cải tiến trên máy chủ BTC',True)
    for text in [
        'Nhóm ưu tiên điểm máy chủ cho các bài Practice. Đợt mới dùng 8 bản đồ khác nhau: Q01, Q03, Q04, Q06, Q07, Q08, Q09 và Q10; Q02 trùng cấu hình Q01 nên không được tính thành mẫu độc lập, còn Q05 chưa có trong danh sách ID đã khảo sát. Trước mỗi lần reset Practice của đội 62, nhóm sao lưu kết quả, trạng thái, kế hoạch và replay hiện có. Các thao tác reset và nộp chỉ thực hiện trên lượt tự luyện riêng của đội mình.',
        'Đội hình tour được lấy từ experiments/tour_results/pipeline.json, vốn chọn offline với ngân sách 20 giây. Khi nộp, nhóm giữ nguyên đội hình và dành 5 giây/ngày cho planner; đầu mỗi ngày, runner lấy lại vị trí, nhiên liệu và traffics từ API BTC. Điểm cuối trận được lấy từ /api/game/result. Sau đó, cùng output đã nộp được chạy lại trong simulator và đối chiếu từng frame bằng experiments/compare_api.py. Bảng kết quả chỉ được lập sau khi có đủ phản hồi máy chủ và đối chiếu replay.',
        'Đối chứng Greedy trong bảng mới là điểm BTC đã chấm ở đợt nộp 9 bài trước đó, với bước chọn vai trò 8 giây và tối đa 20 giây/ngày. Vì cách chọn đội hình và ngân sách khác nhau, đây là so sánh hai quy trình hoàn chỉnh, không phải phép tách riêng tác dụng của tìm kiếm cục bộ hay ô neo. Các đối chứng cùng ngân sách và các phản ví dụ ban đầu vẫn được giữ ở mục 6.1–6.8. Điểm đội 57 là ảnh chụp kết quả máy chủ đã đọc bằng tài khoản được chủ dự án cung cấp; nhóm không biết thuật toán và thời gian tính của họ, nên chỉ dùng làm mốc tham khảo.',
        'Mọi điểm tour trong bảng chính của mục 6.9 là điểm BTC, không phải thay thế bằng điểm offline cao hơn. Kết quả tự sinh trong pipeline được đo trước lần bổ sung bảo vệ D và chỉ là thí nghiệm bổ trợ; không gộp chúng thành bằng chứng tổng quát cho phiên bản mới. Nhóm không khẳng định các lượt Practice tự luyện chứng minh chương trình luôn nộp kịp hạn 30 giây trong trận tính giờ.'
    ]:paragraph(a,text)

    a=find('7. Đánh giá giới hạn')
    paragraph(a,'6.9. Kết quả lập tuyến cả đội được Ban tổ chức chấm',True)
    paragraph(a,'Bảng 29 trình bày điểm máy chủ của hai quy trình trên cùng các bản đồ. Cận U và D được tính trực tiếp từ số loại trên bản đồ và số ngày. Khi một output hợp lệ đạt hai cận này, U và D đã tối ưu trên bài tương ứng; điều đó chưa xác định giá trị tối ưu của Q.')
    paragraph(a,'Bảng 29. Điểm BTC của đối chứng Greedy và phương pháp lập tuyến cả đội',caption=True)
    rows=[]
    for r in data:
        rows.append([r['case'],str(tuple(baseline[r['case']]['ours_score'])),str(tuple(r['server_score'])),str(tuple(r['ud_bound']))])
    table(a,['Bài','Greedy cũ: BTC chấm (U, D, Q)','Lập tuyến cả đội: BTC chấm (U, D, Q)','Cận trên (U, D)'],rows,[1.2,5,5.5,3.8])
    paragraph(a,f'Phương pháp mới đạt cận U và D trên {ud}/8 bài. So với đối chứng Greedy, cải thiện D cho thấy việc giữ khả năng hoạt động qua nhiều ngày quan trọng hơn chỉ chọn loại mới ở thời điểm hiện tại. Q tiếp tục tăng sau khi đã phủ đủ loại. Kết luận này thuộc các bài đã chạy; nó không chứng minh tour luôn phủ đủ loại trên mọi bản đồ hoặc mọi lần chạy có hạn giờ.')
    picture=a.insert_paragraph_before()
    picture.alignment=WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next=True
    picture.add_run().add_picture(str(ROOT/'docs/figures/tour_server_q01.png'),width=Cm(16))
    paragraph(a,'Hình 9. Diễn biến Q01 từ replay BTC của Greedy và phương pháp lập tuyến cả đội.',caption=True)
    paragraph(a,'Ở Q01, Greedy thu đủ 20 loại trong ba ngày đầu nhưng giảm dần khi nhiên liệu cạn; sau đó các xe gần như chỉ lấy udon tại vị trí còn đứng được. Phương pháp mới dùng một xe tiếp tế và giữ đủ 20 loại ở cả mười ngày. Đồ thị tích lũy số suất vì thế tiếp tục tăng đến cuối trận. Ví dụ này minh họa cơ chế cải thiện của cả quy trình; do cùng lúc thay đội hình, cách dựng tuyến và chuẩn bị ngày sau, nó chưa tách được đóng góp riêng của từng thành phần.')
    paragraph(a,'Bảng 30. Đối chiếu số suất khi hai bên cùng đạt cận U và D',caption=True)
    rows=[]
    for r in data:
        peer=baseline[r['case']]['peer_score'];new=r['server_score']
        cap=pipeline[r['case']]['new']['bound_for_types'][2]
        rows.append([r['case'],str(new[2]),str(peer[2]),f'{new[2]-peer[2]:+d}',str(cap),f'{100*new[2]/cap:.1f}%'])
    table(a,['Bài','Q của tour: BTC','Q đội 57: BTC','Chênh lệch Q','Cận Q theo kho và số xe tuần tra','Q/cận sơ bộ'],rows,[1,2.6,2.6,2.1,4.5,2.7])
    paragraph(a,f'Trong ảnh chụp điểm đối chiếu, phương pháp mới có bộ điểm cao hơn đội 57 ở {wins}/8 bài. Đây là so sánh các lượt nộp cụ thể trên cùng input, không phải kết luận vượt mọi đội hay biết kỷ lục của cuộc thi. Cận Q trong bảng bằng số ngày nhân tổng min(kho tại điểm, số xe tuần tra); nó chưa xét lộ trình, step hay nhiên liệu và không bảo đảm đạt được. Tỷ lệ Q/cận chỉ mô tả mức thu so với nguồn cung được nới lỏng, không phải tỷ lệ so với nghiệm tối ưu. Kết quả offline ban đầu vẫn được lưu riêng trong pipeline.json; không dùng các điểm đó thay điểm BTC trong bảng mới.')
    paragraph(a,'Bảng 31. Điều kiện chạy và kiểm chứng replay cho phương pháp mới',caption=True)
    rows=[]
    for r in data:
        rows.append([r['case'],str(r['types']),f"{r['max_planning_seconds']:.2f}",str(r['fallback_days']),str(r['checked_frames']),str(r['mismatch_count'])])
    table(a,['Bài','Vai trò theo thứ tự xe (0 tuần tra; 1 tiếp tế)','Thời gian lập kế hoạch lớn nhất/ngày (giây)','Số ngày dự phòng','Frame đối chiếu','Sai khác'],rows,[1,4.9,3.2,1.9,2.3,2.2])
    paragraph(a,f'Tám trận có tổng 80 kế hoạch ngày được máy chủ nhận, {frames} frame và {values} giá trị được kiểm tra. Số frame và sai khác chỉ nói về các trường do bộ đối chiếu đo: vị trí, nhiên liệu, thu hàng và trạng thái đường kế tiếp; chúng không phải chứng minh hình thức cho toàn bộ máy chủ. Việc một lần chạy đạt D tối đa cũng không thay thế thử nghiệm nhiều lần để đo độ ổn định của tìm kiếm có giới hạn thời gian.')

    replace('Các đối chứng cho thấy A* có lợi',
        f'Các đối chứng cho thấy A* có lợi trong lượng tìm kiếm của bản triển khai, còn Lookahead hai điểm chưa cải thiện ổn định so với Greedy ở ngân sách ngắn. Phương pháp lập tuyến cả đội mở rộng phạm vi quyết định từ một đích sang tập tuyến có tiếp tế và vị trí cuối ngày. Điểm BTC xác nhận phương pháp đạt cận U và D trên {ud}/8 bản đồ Practice đã thử. Đây là kết quả ứng dụng chính của phiên bản mới; Greedy và Lookahead được giữ để giải thích giới hạn của các thiết kế ban đầu.')
    replace('Các bản đồ BTC đã thử chủ yếu',
        'Các bản đồ BTC đã thử có số mẫu độc lập ít; Q01/Q02 trùng cấu hình. Lượt nộp tour mới dùng đội hình đã chọn offline, nên chưa đánh giá trọn bước chọn vai trò dưới đồng hồ thi thật. Phép so với Greedy thay đổi cả ngân sách và nhiều thành phần cùng lúc; cần thí nghiệm tách cover_types, ô neo và tiếp tế để lượng hóa đóng góp riêng. Đạt cận U, D không chứng minh tối ưu Q, còn cận tiếp cận vẫn lỏng do bỏ qua một phần lịch trình và tài nguyên.')
    replace('Các phép đo ở mục 6.7 và 6.8',
        'Phiên bản mới đã triển khai tìm kiếm cục bộ phá và dựng lại, nhiều kiểu đường Dijkstra và cơ chế ô neo cho ngày sau. Hướng tiếp theo là đo tách từng thành phần ở cùng ngân sách, lặp lại nhiều seed và kiểm tra trên trận tính giờ thật. Nhóm cũng có thể thử lịch tiếp tế nhiều điểm và cách ước lượng giao thông ngày mai tốt hơn; hiện chưa cài 2-opt, 3-opt hay beam search. Mỗi cải tiến cần có đối chứng riêng thay vì chỉ bổ sung tên thuật toán.')
    replace('Đề tài cho thấy một bài toán có nguồn gốc',
        'Đề tài cho thấy kiến thức đồ thị, tham lam, tìm kiếm cục bộ và mô hình tài nguyên có thể phối hợp thành một hệ thống lập kế hoạch đạt kết quả kiểm chứng được. Phương pháp lập tuyến cả đội là bước cải tiến chính sau các đối chứng Greedy và Lookahead; điểm BTC được dùng để xác nhận thành tích trên Practice, còn quy hoạch động cung cấp mốc tối ưu ở bài thu nhỏ. Báo cáo phân biệt rõ phần đã đạt cận, phần vẫn là heuristic và các giới hạn thực nghiệm, thay vì coi một bộ điểm cao là nghiệm tối ưu toàn bộ bài toán.')
    # Ghi nguồn có thể mở kiểm tra thủ công, không đưa token vào báo cáo.
    a=find('Phụ lục A.')
    paragraph(a,'[9] Nhóm thực hiện. Kết quả tour được BTC chấm ngày 20/10/2026: experiments/tour_api_evidence/summary.json; mã tái hiện experiments/verify_tour_server.py; cấu hình và đối chứng experiments/tour_results/server_baseline.json. Có thể xem thủ công điểm hiện tại của tài khoản phenikaa_2 trên website BTC; điểm có thể thay đổi khi người dùng nộp lại.')
    a=find('Phụ lục C.')
    paragraph(a,'B.6. Tái hiện phương pháp lập tuyến cả đội',True)
    paragraph(a,'Để dùng phương pháp mới, giảng viên chạy python -X utf8 play_match.py --algorithm tour --mode practice --seconds 5 ID_TRAN. Chương trình chọn vai trò theo tour nếu trận chưa chọn xe. Để tái hiện đúng đội hình của bảng 31, giảng viên đọc trường types trong summary.json và dùng tùy chọn --types; cần đặt ID trận trước --types để danh sách số không nhận nhầm ID. Ví dụ: python -X utf8 play_match.py ID_TRAN --algorithm tour --mode practice --seconds 5 --types 1 0 0 0. Lệnh này có nộp bài; trận đã kết thúc cần reset thủ công trước.')
    paragraph(a,'Script experiments/verify_tour_server.py dành cho tái hiện phép kiểm chứng: nó có reset Practice riêng đội 62, lưu lịch sử trước reset, nộp 10 ngày rồi đối chiếu replay. Script bỏ qua bài đã có trong summary.json để tránh vô tình ghi đè các kết quả đã báo cáo. Thư mục bằng chứng chi tiết và token được giữ tại máy; chỉ bản tổng hợp không chứa token được đưa lên Git. RL trong phụ lục C là phép mở rộng riêng, không được dùng để giải thích thành tích tour.')
    doc.save(path)
    print('Đã cập nhật phương pháp, thực nghiệm, kết quả, kết luận và hướng dẫn tour.')


if __name__ == '__main__':
    update_report()
