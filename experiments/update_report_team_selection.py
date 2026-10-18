"""Cập nhật trực tiếp báo cáo Word với thuật toán và số liệu chọn đội hình."""

import json
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from hexudon.bounds import score_bounds
from hexudon.model import Board
path = ROOT / "output/docx/bao_cao_hexudon.docx"
rows = json.loads((ROOT / "experiments/team_selection_results/summary.json").read_text(encoding="utf-8"))
server = json.loads((ROOT / "experiments/team_selection_api_evidence/summary.json").read_text(encoding="utf-8"))
doc = Document(path)
backup = ROOT / "scratch/report_before_team_selection.docx"
backup.parent.mkdir(parents=True, exist_ok=True)
if not backup.exists():
    backup.write_bytes(path.read_bytes())
doc = Document(backup)


def insert(anchor, text, style=None):
    return anchor.insert_paragraph_before(text, style=style)


def table_before(anchor, headers, values):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E6EDF1")
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in values:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = str(value)
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    anchor._p.addprevious(table._tbl)
    gap = anchor.insert_paragraph_before("")
    gap.paragraph_format.line_spacing = Pt(3)
    gap.paragraph_format.space_after = Pt(3)
    if len(headers) == 7:
        # Giữ bảng trong bề rộng trang, tránh để Word tự kéo giãn bảy cột.
        table.autofit = False
        widths = [1.15, 2.35, 2.35, 2.35, 2.35, 2.35, 2.35]
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)


bound_cache = {}


def detailed_bounds(case_name, kinds):
    key = case_name, tuple(kinds)
    if key in bound_cache:
        return bound_cache[key]
    config = json.loads((ROOT / f"experiments/team_selection_results/{case_name}.json").read_text(encoding="utf-8"))["config"]
    result = score_bounds(Board.from_dict(config), kinds)
    bound_cache[key] = result
    return result


def upper_bounds(case_name, kinds):
    result = detailed_bounds(case_name, kinds)
    return result["rules"], result["reachability"]


def compact_score(score):
    return "(" + ",".join(map(str, score)) + ")"


bound_rows = []
for row in rows:
    rules, selected = upper_bounds(row["case"], row["types"])
    assert all(a <= b for a, b in zip(row["scores"]["selected"], selected))
    assert all(a <= b for a, b in zip(selected, rules))
    bound_rows.append({"case": row["case"], "rules": rules, "selected_types": selected,
                       "patrol_count": detailed_bounds(row["case"], row["types"])["patrol_count"]})
(ROOT / "experiments/team_selection_results/upper_bounds.json").write_text(
    json.dumps(bound_rows, ensure_ascii=False, indent=2), encoding="utf-8")


paragraphs = list(doc.paragraphs)
for p in paragraphs:
    if p.text.startswith("Bộ chọn cơ cấu thử tối đa ba phương án"):
        p.text = (
            "Chúng em giữ nguyên các ô xuất phát do BTC cung cấp và chỉ lựa chọn vai trò của từng xe. "
            "Mỗi vai trò được mã hóa bằng 0 cho xe tuần tra và 1 cho xe tiếp tế. Với A xe, chương trình "
            "có 2^A − 1 cách phân công còn ít nhất một xe tuần tra. Bộ chọn liệt kê những cách phân công "
            "này, giữ cấu hình toàn tuần tra, mọi cách chọn đúng một xe tiếp tế và hai cấu hình tiếp tế "
            "theo vị trí của bộ chọn cũ. Khi còn chỗ, chương trình bổ sung các ứng viên nhiều xe tiếp tế "
            "theo số xe tiếp tế và khoảng cách lục giác tới vùng các điểm udon. Khoảng cách này chỉ phục vụ "
            "sàng lọc, chưa tính được đầy đủ ao và đường vòng. Danh sách mặc định giữ tối đa 16 ứng viên; "
            "những cấu hình bắt buộc vẫn được giữ khi người dùng đặt giới hạn nhỏ hơn.")
    elif p.text.startswith("Đây là quá trình thử một số cấu hình tiêu biểu"):
        p.text = (
            "Mỗi ứng viên được mô phỏng cả trận bằng cùng một bộ lập kế hoạch. Xe, nhiên liệu, kho "
            "và lịch sử giao thông đều được khởi tạo riêng; giao thông ngày sau được tính từ hoạt động "
            "của chính đội trong hai ngày trước, không giả định mọi ngày đều thông thoáng. Sau lượt sàng "
            "lọc, bộ chọn đánh giá lại tối đa năm đội hình với ngân sách mỗi ngày lớn hơn. Vòng này giữ "
            "cấu hình toàn tuần tra, cấu hình xe cuối tiếp tế, ứng viên thắng trong bộ ba cấu hình cũ và "
            "các ứng viên tốt của lượt đầu. Chương trình so bộ điểm cuối trận theo thứ tự số loại toàn "
            "trận, tổng loại theo ngày và tổng suất; trận chưa mô phỏng xong không được dùng để chọn. "
            "Nếu không đủ thời gian hoàn thành bất kỳ ứng viên nào, bộ chọn trả cấu hình toàn tuần tra. "
            "Đây là tìm kiếm giới hạn thời gian trên một chính sách điều khiển cụ thể, không phải chứng "
            "minh đội hình tối ưu toàn cục.")
    elif "Bộ kiểm thử hiện có 35 phương thức" in p.text:
        p.text = p.text.replace("35 phương thức", "44 phương thức") + (
            " Bốn kiểm thử bổ sung kiểm tra việc duyệt đủ phân công ở nhóm xe nhỏ, giữ nguyên vị trí "
            "xuất phát, chọn theo điểm trận hoàn chỉnh và xử lý khi ngân sách đã hết. Năm kiểm thử "
            "khác kiểm tra cận tiếp cận: nhiên liệu dùng chung qua các ngày, chờ để thu lại, thời "
            "gian tích lũy, tiếp tế trong vùng liên thông và duyệt các kế hoạch nhỏ hợp lệ.")
    elif "Bộ chọn loại xe và điểm hẹn tiếp tế cũng mới xét" in p.text:
        p.text = p.text.replace(
            "Bộ chọn loại xe và điểm hẹn tiếp tế cũng mới xét một số phương án đơn giản, chưa tối ưu phân vùng hoặc lựa chọn thời điểm gặp trên đường.",
            "Bộ chọn loại xe đã được mở rộng bằng sàng lọc và mô phỏng lại, nhưng vẫn bị giới hạn bởi "
            "ngân sách tính toán và bộ điều khiển dùng để đánh giá. Lịch tiếp tế còn dùng trạm hẹn "
            "chung, chưa tối ưu phân vùng hoặc thời điểm gặp trên đường.")

anchor = next(p for p in doc.paragraphs if p.text.startswith("8.4 Quy trình"))
insert(anchor, "Thí nghiệm bổ sung về lựa chọn đội hình", "Heading 2")
insert(anchor,
       "Chúng em thực hiện phép đo riêng trên chín cấu hình BTC và sáu bản đồ tự sinh, gồm hai seed "
       "2026 và 2027 ở mỗi kích thước 8×8, 16×16 và 32×32. Q01 và Q02 có input giống nhau nên không "
       "được xem là hai trường hợp độc lập. Thời hạn chọn đội hình là 12 giây với cấu hình BTC và "
       "4 giây với bản đồ tự sinh. Lượt đánh giá lại độc lập cấp 0,15 giây lập kế hoạch cho mỗi ngày, "
       "dùng Greedy của nhánh ttud_clean và simulator một đội cho mọi cấu hình. Các đối chứng gồm "
       "toàn tuần tra, xe cuối tiếp tế, kết quả chọn trong ba cấu hình cũ và đội hình được bộ chọn mới "
       "đề xuất. Điểm dưới đây khác lượt đánh giá code nhánh Procon2026 vì bộ lập kế hoạch và ngân "
       "sách tính toán khác nhau; chúng em không quy toàn bộ chênh lệch giữa hai nhánh cho việc chọn xe.")
insert(anchor,
       "Ngân sách ngày ở vòng đánh giá độc lập được tách khỏi ngân sách sàng lọc. Vì bộ tìm đường "
       "dùng hạn giờ thực, cùng một input vẫn có thể cho tuyến khác nhau khi tiến trình bị ngắt ở "
       "thời điểm khác nhau. Phép đo giữ lại cả những trường hợp điểm sau đánh giá lại giảm, thay vì "
       "chọn số liệu thuận lợi. Các cấu hình và toàn bộ mảng lệnh từng ngày được lưu trong thư mục "
       "experiments/team_selection_results để có thể chạy lại simulator mà không phụ thuộc tốc độ máy.")

anchor = next(p for p in doc.paragraphs
              if p.style.name == "Heading 1" and p.text.startswith("10 Hướng dẫn"))
insert(anchor, "9.5 Kết quả lựa chọn đội hình đầu trận", "Heading 2")
insert(anchor,
       "Bảng sau trình bày kết quả đánh giá lại bằng simulator, không phải điểm BTC của toàn bộ "
       "chín bài. Danh sách vai trò giữ đúng thứ tự agent trong input; các vị trí không được hoán đổi. "
       "Mỗi bộ điểm có dạng (U, D, Q) và được so theo thứ tự ưu tiên, không cộng ba thành phần.")
insert(anchor,
       "Hai cột cận trên giúp người đọc phân biệt giới hạn nguồn lực với kết quả thuật toán. "
       "Gọi B là số loại udon, H là số ngày, A là tổng số xe, P là số xe tuần tra trong đội hình "
       "đã chọn và s_i là tồn kho được cấp mỗi ngày tại điểm i. Cận theo đề bài được tính bằng "
       "(B, H×B, H×tổng min(s_i,A)); do input hợp lệ có s_i không vượt A, thành phần cuối cũng bằng "
       "H nhân tổng tồn kho các điểm. Cận chỉ xét số xe tuần tra thay A bằng P; nếu P bằng 0 thì cả "
       "ba thành phần bằng 0. Trong một ngày, các cận thô tương ứng là B loại và tổng min(s_i,P) suất. "
       "Cách tính dựa vào việc mỗi xe tuần tra chỉ lấy tối đa một suất tại mỗi điểm trong ngày và "
       "kho được bổ sung lại đầu ngày. Cột cận đội hình áp dụng cho phương án đã chọn, không áp "
       "cho hai cấu hình đối chứng trong cùng hàng.")
insert(anchor,
       "Cột cuối của bảng sử dụng cận tiếp cận theo nhiên liệu, chặt hơn hoặc bằng cận chỉ đếm "
       "xe tuần tra. Với từng xe, chúng em chạy Dijkstra hai lần từ ô xuất phát: một lần tối thiểu "
       "hóa nhiên liệu và một lần tối thiểu hóa thời gian trong điều kiện đường luôn thông thoáng. "
       "Một điểm chỉ được tính cho xe ở ngày d nếu thời gian tối thiểu tới điểm đó không vượt tổng "
       "step của các ngày 1 đến d. Nếu xe không thể gặp tiếp tế trong vùng liên thông của mình, "
       "chi phí nhiên liệu tối thiểu tới điểm còn phải không vượt bình ban đầu; nhiên liệu không "
       "được khởi tạo lại khi sang ngày mới. Hai phép tìm đường được xét độc lập, nên có thể cho "
       "hai tuyến khác nhau; phép tính chỉ kiểm tra điều kiện cần, không khẳng định có một tuyến "
       "đồng thời thỏa mọi ràng buộc.")
insert(anchor,
       "Nếu vùng liên thông có xe tiếp tế, chúng em nới lỏng bằng cách cho phép xe tuần tra nạp "
       "nhiên liệu tại mọi ô trong vùng ngay từ đầu trận. Xe tiếp tế không bị giới hạn nhiên liệu "
       "nhưng thực tế vẫn mất thời gian di chuyển và phải gặp đúng xe; các ràng buộc đó được bỏ "
       "qua trong phép tính cận. Việc bỏ qua chỉ làm tập điểm có thể tiếp cận rộng hơn, nên không "
       "loại nhầm kế hoạch hợp lệ. Xe tiếp tế ở vùng bị ao chia tách không được dùng để nới lỏng "
       "nhiên liệu cho một xe tuần tra thuộc vùng khác.")
insert(anchor,
       "Trong từng ngày, chương trình đếm số loại thuộc hợp các điểm còn có thể tiếp cận và "
       "cộng min(tồn kho điểm, số xe tuần tra đủ điều kiện tới điểm) để chặn số suất. Cận toàn "
       "trận gồm số loại trong hợp cuối trận, tổng cận loại của từng ngày và tổng cận suất của "
       "từng ngày. Mọi điểm thực sự thu được đều phải thỏa các điều kiện cần này, nên kế hoạch "
       "hợp lệ không vượt cận theo từng thành phần. Tuy nhiên, một xe đủ nhiên liệu tới từng "
       "điểm khi xét riêng chưa chắc ghé được tất cả chúng trong cùng hành trình. Cận không "
       "tính lại đầy nhiên liệu mỗi ngày; xe đứng tại một điểm qua nhiều ngày vẫn được tính "
       "khả năng thu tiếp, kể cả khi đã hết nhiên liệu.")
insert(anchor,
       "Trên chín cấu hình BTC và sáu bản đồ tự sinh đã đo, cận tiếp cận chưa giảm so với cận "
       "chỉ xét số xe tuần tra. Trong các cấu hình này, mỗi điểm riêng lẻ đều qua được điều kiện "
       "tiếp cận của đủ số xe cần thiết; khó khăn về ghé nhiều điểm trong một hành trình chưa "
       "được mô hình cận nắm bắt. Điều này không chứng minh thuật toán gần tối ưu. Ngược lại, "
       "phép thử nhỏ trên đường năm ô, một xe có một đơn vị nhiên liệu, hai điểm ở ô 1 và ô 4 "
       "và ba ngày cho cận giảm từ (2,6,6) xuống (1,3,3). Chúng em kiểm tra 456 ngày thuộc các "
       "kế hoạch đối chứng đã lưu và không có điểm thực nghiệm nào vượt cận tương ứng. Chi phí "
       "chính là hai lượt Dijkstra cho mỗi xe tuần tra và tìm vùng đi được của xe tiếp tế, phù "
       "hợp với bản đồ tối đa 32×32. Các kết quả chi tiết từng ngày nằm trong reachability_bounds.json.")
table_before(anchor,
             ["Bài", "Vai trò đã chọn", "Toàn tuần tra", "Xe cuối tiếp tế", "Đội hình chọn",
              "Cận nguồn cung đề bài", "Cận tiếp cận nhiên liệu"],
             [[r["case"], str(r["types"]).replace(" ", ""), compact_score(r["scores"]["all_patrol"]),
               compact_score(r["scores"]["last_supply"]), compact_score(r["scores"]["selected"]),
               *[compact_score(bound) for bound in upper_bounds(r["case"], r["types"])]]
              for r in rows if r["case"].startswith("Q")])

def counts(group, baseline):
    wins = sum(tuple(r["scores"]["selected"]) > tuple(r["scores"][baseline]) for r in group)
    losses = sum(tuple(r["scores"]["selected"]) < tuple(r["scores"][baseline]) for r in group)
    return wins, len(group) - wins - losses, losses

table_before(anchor, ["Nhóm", "Đối chứng", "Tốt hơn", "Bằng nhau", "Kém hơn"],
             [[label, baseline_label, *counts(group, baseline)]
              for label, group in [("Cấu hình BTC", [r for r in rows if r["case"].startswith("Q")]),
                                   ("Bản đồ tự sinh", [r for r in rows if not r["case"].startswith("Q")])]
              for baseline, baseline_label in [("all_patrol", "Toàn tuần tra"),
                                               ("last_supply", "Xe cuối tiếp tế"),
                                               ("old_three", "Bộ ba cấu hình cũ")]])
insert(anchor,
       "Kết quả cho thấy không có số xe tiếp tế cố định phù hợp với mọi input. Bộ chọn giữ toàn "
       "tuần tra ở một số bài có nhiên liệu dồi dào, trong khi Q04 chọn xe số 1 và Q06 chọn xe số 0 "
       "làm tiếp tế. Với Q06, đội hình chọn đạt (15, 133, 145), tốt hơn (15, 111, 192) của toàn tuần "
       "tra do tổng loại theo ngày được ưu tiên hơn tổng suất. Tuy vậy, Q03 và Q10 vẫn kém đối chứng "
       "xe cuối tiếp tế trong lượt đánh giá lại. Ở các bản đồ tự sinh cũng có trường hợp giảm điểm. "
       "Những trường hợp này cho thấy ngân sách ngắn và sai khác giữa sàng lọc với điều khiển thực "
       "tế có thể làm đổi thứ hạng đội hình. Bộ chọn chưa bảo đảm luôn tốt hơn mọi cấu hình có sẵn.")
maximum = max(r["selection_seconds"] for r in rows if r["case"].startswith("Q"))
insert(anchor,
       f"Thời gian chọn đội hình lớn nhất trên chín cấu hình BTC là {maximum:.2f} giây trong lượt "
       "đo có hạn 12 giây. Đây là thời gian tính toán offline, chưa gồm tải cấu hình và gửi lựa chọn "
       "qua mạng; kết quả không chứng minh mọi máy hoặc mọi input đều đáp ứng cùng hạn giờ. Lệnh "
       "select mặc định dành 20 giây và để người dùng tự điều chỉnh theo cửa sổ chọn xe của trận. "
       "Khi tăng số xe, bộ chọn chỉ đánh giá danh sách sàng lọc chứ không mô phỏng hết 2^A − 1 cách.")
insert(anchor, "Kiểm chứng hai đội hình trên máy chấm BTC", "Heading 3")
table_before(anchor, ["Bài", "Vai trò", "Điểm BTC", "Frame đối chiếu", "Sai khác",
                      "Cận nguồn cung đề bài", "Cận tiếp cận nhiên liệu"],
             [[r["question"], str(r["types"]), str(tuple(r["server_score"])),
               r["comparison"]["checked_frames"], r["comparison"]["mismatch_count"],
               *[compact_score(bound) for bound in upper_bounds(r["question"], r["types"])]] for r in server])
insert(anchor,
       "Chúng em nộp nguyên kế hoạch đã lưu của Q04 và Q06 vào Practice riêng của đội 62, sau khi "
       "lưu lịch sử trước đó. Cả 20 kế hoạch ngày đều được BTC chấp nhận. Điểm lấy từ API result "
       "khớp điểm offline; đối chiếu 2.807 frame và 44.220 giá trị không phát hiện sai khác. Đây là "
       "kiểm chứng trên hai kế hoạch cụ thể, không phải xác nhận rằng mọi đội hình và mọi trạng thái "
       "của simulator đều đúng. Chúng em chưa nộp toàn bộ ứng viên lên BTC để tìm đội hình tốt nhất "
       "theo máy chủ. Các bằng chứng nằm trong experiments/team_selection_api_evidence, sử dụng "
       "các endpoint Practice, agent-types, day, actions, replay và result đã dẫn ở phần API.")

anchor = next(p for p in doc.paragraphs
              if p.style.name == "Heading 1" and p.text.startswith("11 Đánh giá"))
insert(anchor, "10.4 Chạy bộ chọn đội hình", "Heading 2")
insert(anchor,
       "Người dùng chạy lệnh select một lần trước khi chọn loại xe. Tệp types.json chứa danh sách "
       "vai trò, còn selection.json lưu ứng viên, pha sàng lọc hoặc đánh giá lại, thời gian và mảng "
       "lệnh từng ngày. Những lệnh này không tự sửa đội hình trên website. Sau khi chốt vai trò, "
       "người dùng truyền tệp đó cho plan hoặc demo; từ ngày thứ hai vẫn phải cung cấp trạng thái "
       "đầu ngày do BTC trả về.")
insert(anchor, 'python -X utf8 src/main.py select --config "Đề bài/input_sample.json" --starts 558 998 130 955 --seconds 20 --output results/types.json --selection-log results/selection.json')
insert(anchor, 'python -X utf8 src/main.py plan --config "Đề bài/input_sample.json" --starts 558 998 130 955 --types-file results/types.json --output results/output.txt')
insert(anchor, "python -X utf8 experiments/benchmark_team_selection.py")
insert(anchor,
       "Lệnh bounds nhận cấu hình đầu trận và vai trò đã chọn, rồi xuất cận nguồn cung, cận "
       "theo số xe tuần tra, cận tiếp cận và chi tiết từng ngày. Kết quả không phải dự đoán "
       "điểm đạt được và không nhận trạng thái giữa trận. Lệnh calculate_bounds kiểm tra lại "
       "các kế hoạch đối chứng được lưu trong thí nghiệm đội hình.")
insert(anchor, "python -X utf8 src/main.py bounds --config output/team_selection/Q04/config.json --types-file output/team_selection/Q04/types.json --output results/bounds.json")
insert(anchor, "python -X utf8 experiments/calculate_bounds.py")

doc.save(path)
verified = Document(path)
assert any(p.text == "9.5 Kết quả lựa chọn đội hình đầu trận" for p in verified.paragraphs)
assert any("mẫu số tính lưu lượng bằng 1" in p.text for p in verified.paragraphs)
assert len(verified.tables) == 14
body = list(verified.element.body)
abstract = next(p for p in verified.paragraphs
                if p.style.name == "Heading 1" and p.text == "Tóm tắt")
abstract_index = body.index(abstract._p)
assert not any(node.tag == qn("w:tbl") for node in body[:abstract_index])
assert not any(p.style.name.startswith("Heading") and body.index(p._p) < abstract_index
               for p in verified.paragraphs)
headings = [p.text for p in verified.paragraphs if p.style.name.startswith("Heading")]
assert headings.index("9.4 Kết quả xác minh simulator bằng máy chủ") < headings.index(
    "9.5 Kết quả lựa chọn đội hình đầu trận") < headings.index("10 Hướng dẫn chạy và kiểm chứng")
assert headings.index("10.3 Tái hiện thực nghiệm") < headings.index(
    "10.4 Chạy bộ chọn đội hình") < headings.index("11 Đánh giá giới hạn và hướng phát triển")
print("DOCX updated and reopened; traffic scope preserved; tables:", len(verified.tables))
