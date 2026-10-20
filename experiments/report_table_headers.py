"""Đặt tên cột theo ý nghĩa phép đo, tránh bắt người đọc tự đoán chữ viết tắt."""

import re
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml.ns import qn


HEADERS = [
    ['Quyết định cần giải trong bài toán', 'Thuật toán hoặc phương pháp được sử dụng', 'Tiêu chí đánh giá kết quả'],
    ['Đặc điểm cần so sánh', 'Khả năng của xe tuần tra', 'Khả năng của xe tiếp tế'],
    ['Địa hình tại ô xe rời đi', 'Số step để di chuyển sang ô kề', 'Nhiên liệu xe tuần tra tiêu hao mỗi lần di chuyển'],
    ['Nhóm thông tin đầu vào', 'Tên trường hoặc dữ liệu cần đọc', 'Ý nghĩa và thời điểm sử dụng'],
    ['Phương án minh họa', 'Điểm thu thập: loại toàn trận, loại theo ngày, số suất (U,D,Q)', 'Kết quả so với phương án A có điểm (3,4,5)'],
    ['Thuật toán tìm đường', 'Tiêu chí chọn trạng thái để mở rộng', 'Chất lượng được bảo đảm nếu không hết giờ', 'Mục đích sử dụng trong báo cáo'],
    ['Đặc điểm cần so sánh', 'Greedy: chọn một điểm kế tiếp', 'Lookahead: xét trước hai điểm', 'Quy hoạch động: giải chính xác bài nhỏ'],
    ['Thành phần trong hệ thống', 'Nhiệm vụ khi tạo kế hoạch', 'Giới hạn của thành phần'],
    ['Nội dung cần phân biệt', 'Cách thực hiện ở nhánh huấn luyện gốc', 'Cách thực hiện trong lượt đánh giá của báo cáo'],
    ['Nhóm thí nghiệm', 'Dữ liệu đầu vào và phương pháp so sánh', 'Loại kết quả được sử dụng'],
    ['Kích thước bản đồ / thuật toán', 'Truy vấn tìm được đường / tổng số truy vấn', 'Số truy vấn dừng vì hết giờ', 'Thời gian tìm đường trung bình (ms)', 'Số nhãn trạng thái mở rộng trung bình'],
    ['Kích thước bản đồ / thuật toán', 'Số truy vấn cùng thành công ở cả 5 thuật toán', 'Thời gian di chuyển trung bình (step)', 'Nhiên liệu tiêu hao trung bình', 'Số lần đi sang ô kề trung bình'],
    ['Kích thước bản đồ / chiến lược', 'Số loại udon khác nhau toàn trận (trung bình mỗi trận)', 'Tổng số loại udon cộng qua các ngày (trung bình mỗi trận)', 'Tổng số suất udon thu được (trung bình mỗi trận)'],
    ['Kích thước bản đồ / chiến lược', 'Thời gian lập kế hoạch trung bình mỗi trận (giây)', 'Số trận hợp lệ / tổng số trận', 'Số lần dùng kế hoạch chờ dự phòng'],
    ['Kích thước bản đồ', 'Số cặp Lookahead có điểm cao hơn Greedy', 'Số cặp hai chiến lược có điểm bằng nhau', 'Số cặp Lookahead có điểm thấp hơn Greedy'],
    ['Mã bài thử nghiệm', 'Điểm khi mọi xe đều tuần tra (U,D,Q)', 'Điểm khi xe cuối là xe tiếp tế (U,D,Q)', 'Điểm của đội hình được bộ chọn chọn ra (U,D,Q)'],
    ['Mã bài thử nghiệm', 'Vai trò từng xe theo thứ tự input (0: tuần tra; 1: tiếp tế)', 'Số đội hình trong lượt sàng lọc', 'Thời gian chọn đội hình (giây)'],
    ['Mã bài thử nghiệm', 'Điểm chấm bằng bộ mô phỏng (U,D,Q)', 'Cận trên theo nguồn cung, chưa cố định đội hình', 'Cận trên có xét vị trí, nhiên liệu và đội hình đã chọn'],
    ['Nhóm dữ liệu / cấu hình dùng để đối chiếu', 'Số bài đội hình chọn có điểm cao hơn', 'Số bài hai cấu hình có điểm bằng nhau', 'Số bài đội hình chọn có điểm thấp hơn'],
    ['Mã bài thử nghiệm', 'Điểm máy chủ Ban tổ chức trả về (U,D,Q)', 'Số khung trạng thái được đối chiếu', 'Số sai khác với máy chủ'],
    ['Mã bài thử nghiệm', 'Điểm của Greedy/Lookahead đối chứng (U,D,Q)', 'Điểm của học tăng cường kết hợp Greedy (U,D,Q)'],
    ['Mã bài thử nghiệm', 'Số ngày chọn kế hoạch do mô hình học tăng cường đề xuất', 'Số ngày giữ kế hoạch Greedy', 'Tổng số ngày của trận'],
    ['Kích thước bản đồ', 'Số kế hoạch được kiểm tra', 'Thời gian kiểm tra trung vị (ms)', 'Thời gian kiểm tra lớn nhất (ms)', 'Kế hoạch hợp lệ / tổng số kế hoạch'],
    ['Số xe trong đội', 'Số step được cấp mỗi ngày', 'Thời gian kiểm tra trung vị (ms)', 'Thời gian kiểm tra lớn nhất (ms)'],
    ['Bộ dữ liệu lịch sử trận đấu dùng để đối chiếu', 'Số ngày được kiểm tra', 'Số khung trạng thái được đối chiếu', 'Số giá trị trạng thái được so sánh', 'Số sai khác với máy chủ'],
    ['Kích thước bản đồ / chiến lược', 'Số bài đạt điểm tối ưu / tổng số bài', 'Số loại toàn trận còn thiếu so với tối ưu (trung bình)', 'Trung bình tổng loại theo ngày còn thiếu khi số loại toàn trận đã tối ưu', 'Trung bình số suất còn thiếu khi cả hai tiêu chí loại đã tối ưu'],
    ['Kích thước bản đồ', 'Ngân sách lập kế hoạch mỗi ngày (giây)', 'Số cặp Lookahead thắng Greedy', 'Số cặp có điểm bằng nhau', 'Số cặp Lookahead thua Greedy', 'Số lần gọi tìm đường trung bình mỗi trận (Greedy / Lookahead)'],
    ['Mã lệnh trong output', 'Hướng di chuyển từ ô hiện tại', 'Ký hiệu hướng bằng tiếng Anh'],
    ['Lỗi trong kế hoạch nộp', 'Mã trạng thái máy chủ trả về (HTTP)', 'Cách máy chủ xử lý'],
    ['Nhóm dữ liệu hoặc tài nguyên', 'Có đi kèm kho mã nguồn?', 'Cách kiểm tra hoặc tạo lại dữ liệu'],
]


def apply_readable_headers(doc, clean_run):
    assert len(doc.tables) == len(HEADERS)
    widths = {10: [4, 3.1, 2.5, 3.2, 3.2], 11: [4, 3.3, 2.9, 2.9, 2.9],
              25: [3.3, 2.5, 3.2, 3.5, 3.5], 26: [2.1, 2.6, 2.3, 2.1, 2.3, 4.6]}
    for index, (table, headers) in enumerate(zip(doc.tables, HEADERS)):
        assert len(table.columns) == len(headers)
        for cell, text in zip(table.rows[0].cells, headers):
            paragraph = cell.paragraphs[0]
            size = paragraph.runs[0].font.size.pt
            paragraph.text = text
            for run in paragraph.runs:
                clean_run(run, size)
                run.bold = True
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if index in widths:
            for column, width in zip(table.columns, widths[index]):
                column.width = Cm(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths[index]):
                    cell.width = Cm(width)
        # Hiển thị đủ hai chiều của bản đồ để giá trị 8 không bị hiểu thành số ô tổng cộng.
        if index in (10, 11, 12, 13):
            for row in table.rows[1:]:
                paragraph = row.cells[0].paragraphs[0]
                for run in paragraph.runs:
                    run.text = re.sub(r'^(\d+) / ', r'\1×\1 / ', run.text)
        elif index in (14, 22):
            for row in table.rows[1:]:
                for run in row.cells[0].paragraphs[0].runs:
                    if run.text.isdigit():
                        run.text = run.text + '×' + run.text
        if index == 25:
            for row in table.rows[1:]:
                for cell in row.cells[3:]:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.text = re.sub(r'\(n=(\d+)\)', r'(\1 bài)', run.text)
        if index == 11:
            # Giảm khoảng đệm để bảng không để lại chỉ hai dòng ở trang kế tiếp.
            for row in table.rows:
                for cell in row.cells:
                    for margin in cell._tc.xpath('./w:tcPr/w:tcMar/w:top | ./w:tcPr/w:tcMar/w:bottom'):
                        margin.set(qn('w:w'), '40')
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
        if index == 26:
            # Bảng độ nhạy vừa một trang, nên giữ nguyên khối cùng hàng tiêu đề dài.
            for row_index, row in enumerate(table.rows):
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
